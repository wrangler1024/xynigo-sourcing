# -*- coding: utf-8 -*-
"""退款跟踪导出的列契约：列序、阶段文案、时间格式、空值。

盯的是「导出列与工作台 ④ 表格漂移」这类问题——只数格子个数的断言
挡不住整列错位，所以这里逐列比对表头与取值来源。
"""
from __future__ import annotations

import pathlib
import re
from io import BytesIO

from openpyxl import load_workbook

from xynigo_auth.after_sale_export import (
    CLAIM_HEADERS,
    CLAIM_STATUS_LABELS,
    HEADERS,
    MIME_XLSX,
    build_after_sale_claim_export,
    build_after_sale_track_export,
)


def _row(**overrides) -> dict[str, object]:
    row: dict[str, object] = {
        "refundBillId": "2390833880014851",
        "orderNo": "GSH1RV13Y00NQUV",
        "environmentSerial": "4586",
        "storeName": "合成店铺",
        "status": "ok",
        "phase": "refunded",
        "phaseLabel": "已退款",
        "countdown": "",
        "refundAccount": "****7935",
        "amount": "270.22",
        "goodsImg": "https://img.shein.com/synthetic-goods.jpg",
        "checkedAt": "2026-09-16T01:08:17+00:00",
        "note": None,
        "errorSummary": None,
    }
    row.update(overrides)
    return row


def _sheet(rows):
    content, filename, mime = build_after_sale_track_export(rows)
    workbook = load_workbook(BytesIO(content))
    return workbook, workbook.active, filename, mime


def test_export_header_order_matches_workbench_table() -> None:
    """列序 = 工作台 ④ 表头，逐字钉死（改列必须同时改 UI 与这里）。"""
    assert HEADERS == (
        "环境序号", "订单号", "商品图", "退款单号", "退款信用卡",
        "退款金额", "阶段", "剩余倒计时", "最近检查", "备注",
    )
    sheet = _sheet([])[1]
    assert [cell.value for cell in sheet[1]] == list(HEADERS)
    assert sheet.title == "退款跟踪"


def test_export_row_values_follow_header_order() -> None:
    """一行数据必须按表头顺序落格——整列错位是这块最易犯的错。"""
    _wb, sheet, _filename, mime = _sheet([_row()])
    assert mime == MIME_XLSX
    # 空列落成真正的空单元格（openpyxl 读回来是 None），不写 “—” 占位符，
    # 这样 Excel 里能直接筛选/求和
    assert [cell.value for cell in sheet[2]] == [
        "4586", "GSH1RV13Y00NQUV", "https://img.shein.com/synthetic-goods.jpg",
        "2390833880014851", "****7935", "270.22", "已退款", None,
        "2026-09-16 01:08", None,
    ]
    assert sheet.max_row == 2


def test_export_phase_falls_back_to_chinese_label() -> None:
    """执行器没带 phaseLabel 时按中文表补齐；未知阶段原样导出不吞掉。"""
    _wb, sheet, _f, _m = _sheet([
        _row(phase="reviewing", phaseLabel=""),
        _row(phase="rejected", phaseLabel=""),
        _row(phase="submitted", phaseLabel=""),
        _row(phase="overdue", phaseLabel=""),
        _row(phase="weird_new_phase", phaseLabel=""),
        _row(phase="refunded", phaseLabel="平台已退款"),
    ])
    assert [sheet.cell(row=index, column=7).value
            for index in range(2, 8)] == [
        "审核中", "已拒绝", "已受理", "超期未出结果",
        "weird_new_phase", "平台已退款",
    ]


def test_export_checked_at_and_note_fallbacks() -> None:
    """最近检查统一 `YYYY-MM-DD HH:MM`；备注缺 note 时取 errorSummary。"""
    _wb, sheet, _f, _m = _sheet([
        _row(checkedAt="2026-09-16T01:08:17+00:00"),
        _row(checkedAt=""),
        _row(checkedAt="不是时间"),
        _row(note="拒绝理由：地址不可达"),
        _row(note=None, errorSummary="回访失败：环境占用"),
    ])
    assert [sheet.cell(row=index, column=9).value
            for index in range(2, 5)] == [
        "2026-09-16 01:08", None, "不是时间",
    ]
    assert [sheet.cell(row=index, column=10).value
            for index in range(5, 7)] == [
        "拒绝理由：地址不可达", "回访失败：环境占用",
    ]


def test_export_filename_carries_utc_stamp() -> None:
    content, filename, _mime = build_after_sale_track_export([])
    assert filename.startswith("退款跟踪结果_")
    assert filename.endswith(".xlsx")
    assert len(filename) == len("退款跟踪结果_") + 8 + len(".xlsx")
    assert filename[len("退款跟踪结果_"):len("退款跟踪结果_") + 8].isdigit()
    assert content[:2] == b"PK"  # xlsx 就是 zip 容器


def test_export_without_rows_still_yields_readable_sheet() -> None:
    """没有可回访行时也要给出可用文件（只有表头），不能让页面拿到半截下载。"""
    workbook, sheet, _f, _m = _sheet([])
    assert [cell.value for cell in sheet[1]] == list(HEADERS)
    assert sheet.max_row == 1
    assert workbook.sheetnames == ["退款跟踪"]


# ===== ③ 提交结果（批次）导出 =====
def _claim_row(**overrides) -> dict[str, object]:
    row: dict[str, object] = {
        "environmentSerial": "4589",
        "orderNo": "GSH1RV19M00NEMB",
        "goodsImg": "//img.ltwebstatic.com/v4/j/pi/synthetic.jpg",
        "deliveredAt": "04 Sep 2026 16:56:59",
        "refundBillId": "2390853897109507",
        "refundPath": "Cuenta original de pago",
        "refundAccount": "****0212",
        "status": "ok",
        "submittedAt": "2026-09-16T06:16:50+00:00",
        "note": None,
        "errorSummary": None,
    }
    row.update(overrides)
    return row


def _claim_sheet(rows):
    content, filename, mime = build_after_sale_claim_export(rows)
    return load_workbook(BytesIO(content)).active, filename, mime


def test_claim_export_header_order_matches_workbench_table() -> None:
    """列序 = 工作台 ③ 表头，逐字钉死（改列必须同时改 UI 与这里）。"""
    assert CLAIM_HEADERS == (
        "环境序号", "订单号", "商品图", "售后类型", "送达时间", "退款单号",
        "退款路径", "退款信用卡", "状态", "操作时间", "备注",
    )
    sheet, _filename, mime = _claim_sheet([])
    assert [cell.value for cell in sheet[1]] == list(CLAIM_HEADERS)
    assert mime == MIME_XLSX
    assert sheet.title == "提交结果"


def test_claim_export_row_values_follow_header_order() -> None:
    sheet, filename, _mime = _claim_sheet([_claim_row()])
    assert [cell.value for cell in sheet[2]] == [
        "4589", "GSH1RV19M00NEMB",
        "//img.ltwebstatic.com/v4/j/pi/synthetic.jpg", "丢件退款",
        "04 Sep 2026 16:56:59", "2390853897109507",
        "Cuenta original de pago", "****0212", "已受理 · 退款审核中",
        "2026-09-16 06:16", None,
    ]
    assert filename.startswith("售后提交结果_")
    assert filename.endswith(".xlsx")


def test_claim_export_status_labels_match_web_pills() -> None:
    """状态文案必须与工作台 ③ 的 AS_CLAIM_PILL 同一套（导出与页面不能各说各话）。"""
    web = (pathlib.Path(__file__).parents[3] / "src" / "purchase_tool"
           / "web" / "index.html").read_text(encoding="utf-8")
    block = web[web.index("const AS_CLAIM_PILL = {"):]
    block = block[:block.index("};")]
    pairs = re.findall(r"(\w+):\s*\['\w+',\s*'([^']+)'\]", block)
    assert pairs, "未解析到 Web 的 AS_CLAIM_PILL"
    for status, label in pairs:
        if status in CLAIM_STATUS_LABELS:
            assert CLAIM_STATUS_LABELS[status] == label, (
                f"状态 {status} 的文案与页面不一致："
                f"导出={CLAIM_STATUS_LABELS[status]} 页面={label}")
    # 反向：页面有的状态，导出也得有（否则导出会漏文案）
    assert set(dict(pairs)) - set(CLAIM_STATUS_LABELS) == set()


def test_claim_export_keeps_failure_reason() -> None:
    """失败行必须能看出原因：备注列取 note，缺 note 时落到 errorSummary。"""
    sheet, _filename, _mime = _claim_sheet([
        _claim_row(status="blocked", note="该订单已无可申请售后的包裹"),
        _claim_row(status="login", note=None, errorSummary="登录态失效"),
    ])
    assert sheet.cell(row=2, column=9).value == "不可申请 · 已提交过"
    assert sheet.cell(row=2, column=11).value == "该订单已无可申请售后的包裹"
    assert sheet.cell(row=3, column=9).value == "失败 · 未登录"
    assert sheet.cell(row=3, column=11).value == "登录态失效"


def test_scan_export_preserves_order_numbers_status_and_literal_text():
    from xynigo_auth.after_sale_export import build_after_sale_scan_export, SCAN_HEADERS
    rows = [dict(environmentSerial="0012", storeName="=1+1", orderNo="0001234567890123456789",
                 trackingNo="001234567890123456789", status="ok", claimable=True,
                 amount="0.00", packageCount=2, errorSummary="=HYPERLINK(1)"),
            dict(environmentSerial="0013", orderNo="ORDER2", status="empty", claimable=False),
            dict(environmentSerial="0014", status="fail", errorSummary="读取失败")]
    data, filename, mime = build_after_sale_scan_export(rows)
    sheet = load_workbook(BytesIO(data)).active
    assert tuple(c.value for c in sheet[1]) == SCAN_HEADERS
    assert sheet['A2'].value == '0012'
    assert sheet['C2'].value == '0001234567890123456789'
    assert sheet['I2'].value == '001234567890123456789'
    assert sheet['B2'].value == '=1+1' and sheet['B2'].data_type == 's'
    assert sheet['K2'].data_type == 's'
    assert sheet['G2'].value == 0 and sheet['G2'].data_type == 'n'
    assert sheet['H2'].value == 2
    assert [sheet.cell(i, 10).value for i in range(2, 5)] == ['可申请', '暂不可申请', '失败']
    assert sheet.freeze_panes == 'A2' and sheet.auto_filter.ref == 'A1:K4'
    assert filename.endswith('.xlsx') and mime == MIME_XLSX
    empty, _, _ = build_after_sale_scan_export([])
    assert load_workbook(BytesIO(empty)).active.max_row == 1
