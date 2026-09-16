"""Synthetic roundtrip for product snapshots, completion time and diagnostics."""
import io
import runpy
from pathlib import Path

import pytest
import sqlalchemy as sa
from alembic.migration import MigrationContext
from alembic.operations import Operations
from openpyxl import load_workbook
from pydantic import ValidationError

from xynigo_auth.operation_contract import AfterSaleClaimItem, AfterSaleClaimProgressRow
from test_after_sale_claim import _e2e_setup, _lease_and_start, AS_CAPABILITIES, CLIENT_VERSION
from test_executor_channel import CSRF, device_headers, heartbeat


def test_product_facts_and_action_time_roundtrip_through_progress_history_and_export(tmp_path):
    for web, device, ids, database in _e2e_setup(tmp_path):
        credential=ids['credential']
        heartbeat(device,credential,capabilities=AS_CAPABILITIES,client_version=CLIENT_VERSION)
        products={'goodsImages':['https://img.ltwebstatic.com/synthetic-a.jpg','https://img.ltwebstatic.com/synthetic-b.jpg'],
                  'goodsItems':[{'name':'Synthetic A','specification':'M','quantity':1},{'name':'Synthetic B','quantity':1}],
                  'itemCount':2}
        created=web.post('/v1/operation-runs/after-sale-claim',headers=CSRF,json={
            'executorId':ids['executorId'],'idempotencyKey':'p0-product-snapshot',
            'items':[{'orderNo':'SYNTH','environmentSerial':'ENV',**products}]} )
        assert created.status_code==202,created.text
        run=created.json()['data']['runId']
        task,lease=_lease_and_start(device,credential,expect_type='after.sale.claim.v1')
        # The queue contains no repeated product arrays. The run owns them.
        from xynigo_auth.models import AfterSaleClaimRun
        with database.session_factory() as session:
            import uuid
            saved=session.get(AfterSaleClaimRun,uuid.UUID(run)).request_summary['items'][0]
            assert saved['goodsImages']==products['goodsImages'] and saved['itemCount']==2
        row={'orderNo':'SYNTH','environmentSerial':'ENV','status':'ok','refundBillId':'12345',
             'submittedAt':'2026-09-01T00:00:10+00:00','operationCompletedAt':'2026-09-01T00:00:20+00:00',
             'submissionError':'synthetic missing redirect',
             'refunds':[{'refundBillId':'12345','packageNo':'PKG','source':'recovered_verified',
                         'phase':'reviewing','phaseLabel':'审核中','refundPath':'Cuenta original de pago',
                         'detailsNote':'退款账户未读取'}]}
        def progress(value):
            return device.post(f'/v1/executor-channel/tasks/{task}/progress',headers=device_headers(credential),json={
                'leaseToken':lease,'phase':'after_sale.claim.running','current':1,'total':1,'snapshot':{'rows':[value]}})
        assert progress(row).status_code==200
        row['operationCompletedAt']='2026-09-01T00:01:00+00:00'
        row['refunds'][0].update(refundAccount='****1234',detailsNote='')
        assert progress(row).status_code==200
        for path in [f'/v1/operation-runs/after-sale-claim/{run}',f'/v1/operation-runs/after-sale-claim/history/{run}']:
            response=web.get(path);assert response.status_code==200,response.text
            result=response.json()['data']['rows'][0]
            assert result['goodsImages']==products['goodsImages'] and result['itemCount']==2
            assert result['goodsItems'][0]['specification']=='M'
            assert result['operationCompletedAt'].startswith('2026-09-01T00:00:20')
            assert result['submissionError']=='synthetic missing redirect'
            assert result['refunds'][0]['refundAccount']=='****1234'
            assert not result['refunds'][0]['detailsNote']
        exported=web.get(f'/v1/operation-runs/after-sale-claim/history/{run}/export')
        sheet=load_workbook(io.BytesIO(exported.content)).active
        assert sheet.cell(2,3).value=='\n'.join(products['goodsImages'])
        assert sheet.cell(2,12).value==2 and 'Synthetic A' in sheet.cell(2,13).value
        assert '审核中' in sheet.cell(2,9).value
        assert 'synthetic missing redirect' in sheet.cell(2,11).value
        heartbeat(device,credential,capabilities=[cap for cap in AS_CAPABILITIES if cap!='after.sale.claim-evidence.v1'],client_version=CLIENT_VERSION)
        old=web.post('/v1/operation-runs/after-sale-claim',headers=CSRF,json={
            'executorId':ids['executorId'],'idempotencyKey':'p0-older-executor',
            'items':[{'orderNo':'SYNTHNEXT','environmentSerial':'ENV'}]})
        assert old.status_code==409 and 'executor_after_sale_evidence_upgrade_required' in old.text
        break


@pytest.mark.parametrize('status,completed', [('verifying','2026-09-01T00:00:00Z'),('ok','2026-09-01T00:00:00'),('ok','yesterday')])
def test_completion_requires_terminal_row_and_explicit_timezone(status,completed):
    with pytest.raises(ValidationError):
        AfterSaleClaimProgressRow(orderNo='SYNTH',status=status,operationCompletedAt=completed)


def test_legacy_rows_do_not_invent_completion_or_product_counts():
    old=AfterSaleClaimProgressRow(orderNo='SYNTH',status='ok',submittedAt='2026-09-01T00:00:00Z')
    assert old.operationCompletedAt is None
    item=AfterSaleClaimItem(environmentSerial='ENV',orderNo='SYNTH',goodsImg='one.jpg')
    assert item.itemCount is None and item.goodsImages==[]


def test_export_completion_keeps_seconds_and_unambiguous_timezone():
    from xynigo_auth.after_sale_export import _operation_time_text
    from datetime import datetime
    utc=_operation_time_text('2026-09-01T00:00:20Z')
    local=_operation_time_text('2026-09-01T08:00:20+08:00')
    assert utc=='2026-09-01 00:00:20+00:00'
    assert local=='2026-09-01 08:00:20+08:00'
    assert datetime.fromisoformat(utc)==datetime.fromisoformat(local)


def test_result_facts_migration_preserves_existing_rows():
    migration=runpy.run_path(str(Path(__file__).resolve().parents[1]/'migrations/versions/0042_after_sale_result_facts.py'))
    engine=sa.create_engine('sqlite://')
    with engine.begin() as connection:
        connection.execute(sa.text('CREATE TABLE after_sale_claim_results (id INTEGER PRIMARY KEY, status TEXT)'))
        connection.execute(sa.text("INSERT INTO after_sale_claim_results VALUES (1,'uncertain')"))
        with Operations.context(MigrationContext.configure(connection)):migration['upgrade']()
        row=connection.execute(sa.text('SELECT goods_images,item_count,operation_completed_at FROM after_sale_claim_results')).one()
        assert row==('[]',None,None)
        with Operations.context(MigrationContext.configure(connection)):migration['downgrade']()
        assert connection.execute(sa.text('SELECT status FROM after_sale_claim_results')).scalar()=='uncertain'


def test_new_empty_optional_metadata_preserves_pre_upgrade_idempotency_hash():
    import hashlib,json,uuid
    from xynigo_auth.operation_contract import AfterSaleClaimRunCreateBody
    from xynigo_auth.operation_service import _payload_hash
    body=AfterSaleClaimRunCreateBody(executorId=uuid.UUID(int=1),idempotencyKey='synthetic-p0-key',
        items=[{'environmentSerial':'ENV','orderNo':'SYNTH'}])
    old=body.model_dump(mode='json');old.pop('retryFromRunId')
    for key in ('goodsImages','goodsItems','itemCount'):old['items'][0].pop(key)
    canonical=json.dumps(old,ensure_ascii=False,sort_keys=True,separators=(',',':'))
    assert _payload_hash(body)==hashlib.sha256(canonical.encode()).hexdigest()
