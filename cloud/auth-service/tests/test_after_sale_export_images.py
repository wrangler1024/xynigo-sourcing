"""Actual XLSX packages retain all product images and readable failure rows."""
import io
import zipfile
import threading
from concurrent.futures import ThreadPoolExecutor
from xml.etree import ElementTree as ET

import pytest
from PIL import Image
from openpyxl import load_workbook

from xynigo_auth import after_sale_export_images as pictures
from xynigo_auth.after_sale_export import (
    build_after_sale_scan_export, build_after_sale_claim_export, build_after_sale_track_export,
    AfterSaleExportBusy,
)
from xynigo_auth.after_sale_presentation import display_refund_path

BUILDERS = [build_after_sale_scan_export, build_after_sale_claim_export, build_after_sale_track_export]
A = 'https://img.ltwebstatic.com/synthetic-red.png'
B = 'https://img.ltwebstatic.com/synthetic-blue.png'


def picture(color):
    out=io.BytesIO()
    Image.new('RGB',(60,80),color).save(out,format='PNG')
    return out.getvalue()


@pytest.mark.parametrize('build', BUILDERS)
def test_embeds_all_products_as_in_cell_grid_and_borders_every_cell(build):
    calls=[]
    def fetch(url):
        calls.append(url)
        return picture('red' if url==A else 'blue')
    row={'environmentSerial':'DEMO','orderNo':'SYNTH','goodsImg':B,'goodsImages':[A,B],
         'status':'ok','claimable':True,'amount':'0.00','refundBillId':'12345'}
    data,_,_=build([row,dict(row,orderNo='OTHER')],image_fetcher=fetch)
    assert sorted(calls)==sorted([A,B]), 'deduplicate downloads across products and orders'
    sheet=load_workbook(io.BytesIO(data)).active
    column=4 if build==build_after_sale_scan_export else 3
    assert sheet.max_row==3 and sheet.freeze_panes=='A2'
    assert A in sheet.cell(2,column).comment.text and B in sheet.cell(2,column).comment.text
    for cells in sheet:
        for cell in cells:
            assert all(getattr(cell.border,side).style=='thin' for side in ('left','right','top','bottom'))
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        assert len([name for name in archive.namelist() if name.startswith('xl/media/')])==2
        image=Image.open(io.BytesIO(archive.read('xl/media/image1.jpeg')))
        red,blue=image.getpixel((30,40)),image.getpixel((110,40))
        assert red[0]>200 and red[2]<50 and blue[2]>200 and blue[0]<50
        xml=ET.fromstring(archive.read('xl/worksheets/sheet1.xml'))
        ns={'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
        assert len([cell for cell in xml.findall('.//m:c',ns) if cell.get('vm')])==2


@pytest.mark.parametrize('build', BUILDERS)
def test_missing_and_empty_images_keep_export_usable(build):
    data,_,_=build([{'orderNo':'SYNTH','goodsImages':[A,B],'note':'Original note','errorSummary':'Original reason'}],
                  image_fetcher=lambda _url:b'not an image')
    sheet=load_workbook(io.BytesIO(data)).active
    column=4 if build==build_after_sale_scan_export else 3
    assert A in sheet.cell(2,column).value and B in sheet.cell(2,column).value
    note_column=next(cell.column for cell in sheet[1] if cell.value=='备注')
    assert '2/2 张未取得' in sheet.cell(2,note_column).value
    empty,_,_=build([],image_fetcher=lambda _url:pytest.fail('no download for empty export'))
    empty_sheet=load_workbook(io.BytesIO(empty)).active
    assert empty_sheet.max_row==1 and all(c.border.top.style=='thin' for c in empty_sheet[1])


def test_download_allowlist_budget_and_partial_grid(monkeypatch):
    calls=[]
    monkeypatch.setattr(pictures,'MAX_UNIQUE_IMAGES',2)
    cache=pictures.collect_thumbnails([{'goodsImages':['http://127.0.0.1/private',A,B]}],
        lambda url:calls.append(url) or picture('red'))
    assert calls==[A] and not cache.get(B)
    grid=pictures.product_image_grid([A,B],cache)
    assert grid and Image.open(io.BytesIO(grid[0])).width>100
    monkeypatch.setattr(pictures,'DOWNLOAD_START_WINDOW',0)
    assert not pictures.collect_thumbnails([{'goodsImg':A}],lambda _:pytest.fail('deadline expired'))[A]


def test_failed_downloads_consume_reserved_budget(monkeypatch):
    monkeypatch.setattr(pictures,'MAX_DOWNLOAD_BYTES',2*(pictures.MAX_IMAGE_BYTES+1))
    calls=[]
    def fail(url):
        calls.append(url)
        raise OSError('response interrupted after unknown bytes')
    cache=pictures.collect_thumbnails([{'goodsImages':[A+str(i) for i in range(20)]}],fail)
    assert len(calls)==2 and not any(cache.values())


def test_large_decoded_image_is_rejected_before_loading_pixels(monkeypatch):
    monkeypatch.setattr(pictures,'MAX_IMAGE_PIXELS',100)
    cache=pictures.collect_thumbnails([{'goodsImg':A}],lambda _:picture('red'))
    assert cache[A] is None


def test_concurrent_exports_are_rejected_without_downloading_twice():
    entered,released=threading.Event(),threading.Event()
    def fetch(_):
        entered.set()
        assert released.wait(5)
        return picture('red')
    with ThreadPoolExecutor(max_workers=1) as pool:
        pending=pool.submit(build_after_sale_claim_export,[{'goodsImg':A}],image_fetcher=fetch)
        try:
            assert entered.wait(5)
            with pytest.raises(AfterSaleExportBusy,match='稍后重试'):
                build_after_sale_track_export([],image_fetcher=lambda _:pytest.fail('must not download'))
        finally:
            released.set()
        assert pending.result()[0].startswith(b'PK')
    assert build_after_sale_track_export([])[0].startswith(b'PK'), 'export slot is released'


@pytest.mark.parametrize('value,expected', [
    ('Cuenta original de pago / 其他退款渠道（名称未取得）','Cuenta original de pago'),
    ('Cuenta original de pago ／ 其他退款渠道 ( 名称未取得 )','Cuenta original de pago'),
    ('其他退款渠道（名称未取得）',''),
    ('Cuenta original de pago / Cartera SHEIN','Cuenta original de pago / Cartera SHEIN'),
])
def test_path_display_preserves_known_channels(value,expected):
    assert display_refund_path(value)==expected


def test_supplement_source_does_not_replace_exported_operation_time():
    data,_,_=build_after_sale_claim_export([{
        'orderNo':'SYNTH','operationCompletedAt':'2026-09-01T00:00:20Z',
        'refunds':[{'refundBillId':'12345','refundAccount':'****1234',
            'refundAccountSource':'tracking','refundAccountCheckedAt':'2026-09-02T01:02:03Z'}]}])
    sheet=load_workbook(io.BytesIO(data)).active
    assert sheet['J2'].value=='2026-09-01 00:00:20+00:00'
    assert '最近回访 2026-09-02 01:02:03+00:00' in sheet['K2'].value
