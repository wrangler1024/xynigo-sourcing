# -*- coding: utf-8 -*-
"""售后处理（after.sale.scan.v1 / after.sale.claim.v1）云端全链契约测试。

覆盖：扫描建任务→进度→终态 GET（summary.rows / claimableCount）；
提交建 Run（幂等）→进度行+截图附件→终态快照（行数/终态/截图可读）；
扫描与提交的取消路由把 stop 请求下发到任务；闭集与跨批次白名单 422。
"""
from __future__ import annotations

import base64
import hashlib
import uuid
from datetime import datetime, timedelta, timezone
from io import BytesIO
from urllib.parse import quote

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select

from xynigo_auth.after_sale_export import (
    CLAIM_HEADERS as CLAIM_EXPORT_HEADERS,
)
from xynigo_auth.after_sale_export import HEADERS as EXPORT_HEADERS

from test_executor_channel import (
    CSRF,
    create_pairing_code,
    device_headers,
    heartbeat,
    login,
    pair,
)
from test_purchase_api import build_test_app

AS_CAPABILITIES = [
    "config.read.v1",
    "config.write.v1",
    "workspace.snapshot.v1",
    "after.sale.scan.v1",
    "after.sale.claim.v1",
    "after.sale.track.v1",
    "after.sale.runtime-controls.v1",
    "after.sale.receipt-recovery.v1",
    "after.sale.claim-evidence.v1",
    "after.sale.claim-environment.v1",
    "after.sale.phase-evidence.v1",
    "after.sale.reliable-results.v1",
]
CLIENT_VERSION = "0.17.20"
JPEG = b"\xff\xd8synthetic-after-sale-screenshot\xff\xd9"


def _e2e_setup(tmp_path):
    app, database, _oauth = build_test_app(tmp_path)
    with TestClient(app) as web_client, TestClient(app) as device_client:
        login(web_client)
        paired = pair(
            device_client,
            create_pairing_code(web_client),
            capabilities=AS_CAPABILITIES,
        )
        yield web_client, device_client, {
            "executorId": str(paired["executorId"]),
            "credential": str(paired["deviceCredential"]),
        }, database


def _scan_body(*serials: str) -> dict[str, object]:
    return {
        "idempotencyKey": "as-scan-e2e-00000001",
        "executorId": "",
        "browserMode": "visible",
        "environmentSerials": list(serials or ("4902", "4901")),
    }


def _scan_row(serial: str, *, order_no: str, claimable: bool) -> dict[str, object]:
    """执行器投影后的扫描行（闭集，无 screenshotStatus 等本地字段）。"""
    return {
        "environmentSerial": serial,
        "storeName": "合成店铺-" + serial,
        "accountName": "buyer@example.test",
        "orderNo": order_no,
        "deliveredAt": "2026-09-10",
        "amount": "MXN289.00",
        "status": "ok",
        "claimable": claimable,
        "packageCount": 1 if claimable else 0,
        "trackingNo": "SHEIN" + serial,
        "errorSummary": None,
        "screenshotSha256": None,
    }


def _claim_row(order_no: str, serial: str, *, status: str = "ok") -> dict[str, object]:
    """执行器投影后的提交行（闭集 + 本地会带的 packageCount）。"""
    return {
        "orderNo": order_no,
        "environmentSerial": serial,
        "storeName": "合成店铺-" + serial,
        "status": status,
        "packageNo": "PKG" + order_no[-4:],
        "refundBillId": "RB" + order_no[-4:],
        "refundPath": "原路退回（Cuenta original de pago）",
        "durationSeconds": 42,
        "submittedAt": "2026-09-15T03:00:00+00:00",
        "note": "",
        "errorSummary": None,
        "screenshotSha256": None,
        "packageCount": 1,
    }


def _lease_and_start(device_client, credential, *, expect_type: str) -> tuple[str, str]:
    lease = heartbeat(device_client, credential, capabilities=AS_CAPABILITIES,
                      client_version=CLIENT_VERSION)["task"]
    assert lease is not None and lease["type"] == expect_type, lease
    task_id, lease_token = lease["id"], lease["leaseToken"]
    started = device_client.post(
        f"/v1/executor-channel/tasks/{task_id}/start",
        json={"leaseToken": lease_token},
        headers=device_headers(credential),
    )
    assert started.status_code == 200, started.text
    return task_id, lease_token


# ===== 扫描：建任务 → 进度 → 终态 GET =====
def test_after_sale_scan_progress_and_terminal_summary(tmp_path) -> None:
    for web_client, device_client, ids, database in _e2e_setup(tmp_path):
        executor_id, credential = ids["executorId"], ids["credential"]
        heartbeat(device_client, credential, capabilities=AS_CAPABILITIES,
                  client_version=CLIENT_VERSION)

        created = web_client.post(
            "/v1/after-sale/scan",
            json={**_scan_body("4902", "4901"), "executorId": executor_id},
            headers=CSRF,
        )
        assert created.status_code == 202, created.text
        payload = created.json()["data"]
        assert set(payload) == {"taskId", "status", "executorId"}
        assert payload["status"] == "queued"
        assert payload["executorId"] == executor_id
        task_id = payload["taskId"]

        # 任务表里没有 Run：扫描只是任务 + 任务行上的进度快照
        with database.session_factory() as session:
            from xynigo_auth.models import AfterSaleClaimRun
            assert session.scalar(select(AfterSaleClaimRun)) is None
        # 还没有任何提交批次：latest 返回 null
        latest_empty = web_client.get("/v1/operation-runs/after-sale-claim/latest")
        assert latest_empty.status_code == 200, latest_empty.text
        assert latest_empty.json()["data"] is None

        # 运行中 GET：summary 带上计划环境数（进度条分母）
        running = web_client.get(f"/v1/after-sale/scan/{task_id}")
        assert running.status_code == 200, running.text
        data = running.json()["data"]
        assert data["taskId"] == task_id
        assert data["status"] == "queued"
        assert data["summary"] == {"rows": [], "totalCount": 2,
                                   "claimableCount": 0}
        assert data["lastRuns"] == {}

        lease_task, lease_token = _lease_and_start(
            device_client, credential, expect_type="after.sale.scan.v1")
        assert lease_task == task_id

        # 进度：一条可申请、一条无订单（empty），部分行上报
        progress = device_client.post(
            f"/v1/executor-channel/tasks/{task_id}/progress",
            json={
                "leaseToken": lease_token,
                "phase": "after_sale.scan.running",
                "current": 1,
                "total": 2,
                "snapshot": {"rows": [
                    {**_scan_row("4902", order_no="GSH0001", claimable=True),
                     "platformStatus": "Enviado", "reasonSource": "pre_info",
                     "checkedAt": "2026-09-16T00:00:00+00:00", "reasonCode": ""},
                    {
                        "environmentSerial": "4901",
                        "storeName": "合成店铺-4901",
                        "accountName": "",
                        "orderNo": "",
                        "deliveredAt": "",
                        "amount": "",
                        "status": "empty",
                        "claimable": False,
                        "packageCount": 0,
                        "trackingNo": "",
                        "errorSummary": "该环境订单列表没有可申请售后的已送达订单",
                        "screenshotSha256": None,
                    },
                ]},
            },
            headers=device_headers(credential),
        )
        assert progress.status_code == 200, progress.text

        mid = web_client.get(f"/v1/after-sale/scan/{task_id}").json()["data"]
        assert mid["status"] == "running"
        assert mid["summary"]["totalCount"] == 2
        assert len(mid["summary"]["rows"]) == 2
        assert mid["summary"]["claimableCount"] == 1
        assert mid["summary"]["rows"][0]["platformStatus"] == "Enviado"

        exported = web_client.get(f"/v1/after-sale/scan/{task_id}/export")
        assert exported.status_code == 200, exported.text
        assert exported.headers["x-xynigo-row-count"] == "2"
        assert "no-store" in exported.headers["cache-control"]
        sheet = load_workbook(BytesIO(exported.content)).active
        assert [sheet.cell(i, 1).value for i in (2, 3)] == ["4902", "4901"]
        assert [sheet.cell(i, 10).value for i in (2, 3)] == ["可申请", "无订单"]
        assert web_client.get(f"/v1/after-sale/scan/{task_id}").json()["data"]["status"] == "running"
        assert web_client.get(f"/v1/after-sale/scan/{uuid.uuid4()}/export").status_code == 404
        assert device_client.get(f"/v1/after-sale/scan/{task_id}/export").status_code == 401

        finish = device_client.post(
            f"/v1/executor-channel/tasks/{task_id}/finish",
            json={
                "leaseToken": lease_token,
                "outcome": "succeeded",
                "resultCode": "after_sale_scan_completed",
                "resultSummary": {
                    "rows": [
                        _scan_row("4902", order_no="GSH0001", claimable=True),
                        _scan_row("4901", order_no="GSH0002", claimable=False),
                    ],
                    "totalCount": 2,
                    "claimableCount": 1,
                },
            },
            headers=device_headers(credential),
        )
        assert finish.status_code == 200, finish.text

        final = web_client.get(f"/v1/after-sale/scan/{task_id}")
        assert final.status_code == 200, final.text
        summary = final.json()["data"]["summary"]
        assert final.json()["data"]["status"] == "succeeded"
        assert summary["totalCount"] == 2
        assert summary["claimableCount"] == 1
        assert {row["orderNo"] for row in summary["rows"]} == {"GSH0001", "GSH0002"}
        assert summary["rows"][0]["claimable"] is True
        # 闭集：本地字段不进云端快照
        assert "screenshotStatus" not in summary["rows"][0]
        break


def test_all_orders_scan_finishes_multiple_claimable_orders_in_one_environment(tmp_path) -> None:
    """环境数不是订单数；所有订单扫描同时保留可申请、运输中和核验失败行。"""
    for web_client, device_client, ids, _database in _e2e_setup(tmp_path):
        credential = ids["credential"]
        heartbeat(device_client, credential, capabilities=AS_CAPABILITIES,
                  client_version=CLIENT_VERSION)
        created = web_client.post(
            "/v1/after-sale/scan",
            json={**_scan_body("SYNTHENV"), "executorId": ids["executorId"]},
            headers=CSRF,
        )
        assert created.status_code == 202, created.text
        task_id, lease_token = _lease_and_start(
            device_client, credential, expect_type="after.sale.scan.v1")
        rows = [_scan_row("SYNTHENV", order_no="SYNTH001", claimable=True),
                _scan_row("SYNTHENV", order_no="SYNTH002", claimable=True),
                {**_scan_row("SYNTHENV", order_no="SYNTH003", claimable=False),
                 "status": "skip", "errorSummary": "运输中（Enviado）；当前没有丢件退款申请入口"},
                {**_scan_row("SYNTHENV", order_no="SYNTH004", claimable=False),
                 "status": "fail", "errorSummary": "可申请性核验失败：合成超时"}]
        progress = device_client.post(
            f"/v1/executor-channel/tasks/{task_id}/progress",
            json={"leaseToken": lease_token, "phase": "after_sale.scan.completed",
                  "current": 1, "total": 1, "snapshot": {"rows": rows}},
            headers=device_headers(credential),
        )
        assert progress.status_code == 200, progress.text
        finish_body = {
            "leaseToken": lease_token, "outcome": "succeeded",
            "resultCode": "after_sale_scan_completed",
            "resultSummary": {"rows": rows, "totalCount": 1, "claimableCount": 2},
        }
        # 仍拒绝不实的可申请计数，不能为了多单把计数校验全部放开。
        invalid = {**finish_body, "resultSummary": {
            **finish_body["resultSummary"], "claimableCount": 3}}
        rejected = device_client.post(
            f"/v1/executor-channel/tasks/{task_id}/finish", json=invalid,
            headers=device_headers(credential))
        assert rejected.status_code == 422, rejected.text
        finished = device_client.post(
            f"/v1/executor-channel/tasks/{task_id}/finish", json=finish_body,
            headers=device_headers(credential))
        assert finished.status_code == 200, finished.text
        data = web_client.get(f"/v1/after-sale/scan/{task_id}").json()["data"]
        assert data["status"] == "succeeded"
        assert data["summary"]["totalCount"] == 1
        assert data["summary"]["claimableCount"] == 2
        assert [(r["orderNo"], r["status"], r["errorSummary"]) for r in data["summary"]["rows"]] == [
            (r["orderNo"], r["status"], r["errorSummary"]) for r in rows]
        break


def test_after_sale_scan_rejects_row_extra_field_and_cross_batch_serial(tmp_path) -> None:
    """闭集 extra=forbid 与请求白名单：投影外字段/串批次环境序号必须 422。"""
    for web_client, device_client, ids, _database in _e2e_setup(tmp_path):
        executor_id, credential = ids["executorId"], ids["credential"]
        heartbeat(device_client, credential, capabilities=AS_CAPABILITIES,
                  client_version=CLIENT_VERSION)
        created = web_client.post(
            "/v1/after-sale/scan",
            json={**_scan_body("4902"), "executorId": executor_id},
            headers=CSRF,
        )
        assert created.status_code == 202, created.text
        task_id = created.json()["data"]["taskId"]
        _, lease_token = _lease_and_start(
            device_client, credential, expect_type="after.sale.scan.v1")

        with_extra = device_client.post(
            f"/v1/executor-channel/tasks/{task_id}/progress",
            json={
                "leaseToken": lease_token,
                "phase": "after_sale.scan.running",
                "current": 1,
                "total": 1,
                "snapshot": {"rows": [
                    {**_scan_row("4902", order_no="GSH0001", claimable=True),
                     "screenshotStatus": "ok"},
                ]},
            },
            headers=device_headers(credential),
        )
        assert with_extra.status_code == 422, with_extra.text

        outside = device_client.post(
            f"/v1/executor-channel/tasks/{task_id}/progress",
            json={
                "leaseToken": lease_token,
                "phase": "after_sale.scan.running",
                "current": 1,
                "total": 1,
                "snapshot": {"rows": [
                    _scan_row("9999", order_no="GSH0001", claimable=True),
                ]},
            },
            headers=device_headers(credential),
        )
        assert outside.status_code == 422, outside.text

def test_after_sale_claim_run_progress_finish_and_snapshot(tmp_path) -> None:
    for web_client, device_client, ids, database in _e2e_setup(tmp_path):
        executor_id, credential = ids["executorId"], ids["credential"]
        heartbeat(device_client, credential, capabilities=AS_CAPABILITIES,
                  client_version=CLIENT_VERSION)
        items = [
            {"environmentSerial": "4902", "orderNo": "GSH0001",
             "storeName": "合成店铺-4902", "packageNo": ""},
            {"environmentSerial": "4901", "orderNo": "GSH0002",
             "storeName": "合成店铺-4901", "packageNo": ""},
        ]
        body = {
            "idempotencyKey": "as-claim-e2e-00000001",
            "executorId": executor_id,
            "browserMode": "visible",
            "items": items,
        }
        created = web_client.post(
            "/v1/operation-runs/after-sale-claim", json=body, headers=CSRF)
        assert created.status_code == 202, created.text
        snapshot = created.json()["data"]
        run_id = snapshot["runId"]
        assert snapshot["status"] == "queued"
        assert snapshot["totalCount"] == 2
        assert snapshot["progressTotal"] == 2
        assert snapshot["progressCompleted"] == 0
        assert snapshot["rows"] == []

        # 幂等键重复：同一 Run，不重复下发任务
        replay = web_client.post(
            "/v1/operation-runs/after-sale-claim", json=body, headers=CSRF)
        assert replay.status_code == 202, replay.text
        assert replay.json()["data"]["runId"] == run_id

        task_id, lease_token = _lease_and_start(
            device_client, credential, expect_type="after.sale.claim.v1")

        sha = hashlib.sha256(JPEG).hexdigest()
        progress = device_client.post(
            f"/v1/executor-channel/tasks/{task_id}/progress",
            json={
                "leaseToken": lease_token,
                "phase": "after_sale.claim.running",
                "current": 1,
                "total": 2,
                "snapshot": {
                    "rows": [
                        _claim_row("GSH0001", "4902"),
                        {**_claim_row("GSH0002", "4901", status="running"),
                         "packageNo": "", "refundBillId": "", "refundPath": "",
                         "durationSeconds": None, "submittedAt": None,
                         "packageCount": None},
                    ],
                    "screenshots": [{
                        "orderNo": "GSH0001",
                        "contentType": "image/jpeg",
                        "contentBase64": base64.b64encode(JPEG).decode(),
                        "sha256": sha,
                        "size": len(JPEG),
                    }],
                },
            },
            headers=device_headers(credential),
        )
        assert progress.status_code == 200, progress.text

        partial = web_client.get(
            f"/v1/operation-runs/after-sale-claim/{run_id}").json()["data"]
        assert partial["status"] == "running"
        assert len(partial["rows"]) == 2
        first = [r for r in partial["rows"] if r["orderNo"] == "GSH0001"][0]
        assert first["status"] == "ok"
        assert first["submittedAt"].startswith("2026-09-15T03:00:00")
        assert first["screenshotSha256"] == sha

        # 收尾 progress：执行器在批次结束后会再报一次全终态行（终态回执不带行）
        final_progress = device_client.post(
            f"/v1/executor-channel/tasks/{task_id}/progress",
            json={
                "leaseToken": lease_token,
                "phase": "after_sale.claim.completed",
                "current": 2,
                "total": 2,
                "snapshot": {"rows": [
                    _claim_row("GSH0001", "4902"),
                    _claim_row("GSH0002", "4901"),
                ]},
            },
            headers=device_headers(credential),
        )
        assert final_progress.status_code == 200, final_progress.text

        finish = device_client.post(
            f"/v1/executor-channel/tasks/{task_id}/finish",
            json={
                "leaseToken": lease_token,
                "outcome": "succeeded",
                "resultCode": "after_sale_completed",
                "resultSummary": {
                    "runStatus": "completed",
                    "phase": "after_sale.completed",
                    "progressCompleted": 2,
                    "progressTotal": 2,
                    "totalCount": 2,
                    "successCount": 2,
                    "failedCount": 0,
                    "skippedCount": 0,
                    "stoppedCount": 0,
                    "errorCode": "",
                    "errorSummary": "",
                },
            },
            headers=device_headers(credential),
        )
        assert finish.status_code == 200, finish.text

        final = web_client.get(
            f"/v1/operation-runs/after-sale-claim/{run_id}").json()["data"]
        assert final["status"] == "completed"
        assert final["phase"] == "after_sale.completed"
        assert final["successCount"] == 2
        assert final["failedCount"] == 0
        assert final["skippedCount"] == 0
        assert final["stoppedCount"] == 0
        assert final["progressCompleted"] == 2
        assert final["progressTotal"] == 2
        assert final["stopRequested"] is False
        assert len(final["rows"]) == 2
        assert {row["status"] for row in final["rows"]} == {"ok"}
        assert final["rows"][0]["orderNo"] == "GSH0001"

        # latest 返回同一批次
        latest = web_client.get("/v1/operation-runs/after-sale-claim/latest")
        assert latest.status_code == 200, latest.text
        assert latest.json()["data"]["runId"] == run_id

        # 扫描结果页的 lastRuns：按店铺合并「上次提交到哪一步」
        scan_created = web_client.post(
            "/v1/after-sale/scan",
            json={
                "idempotencyKey": "as-scan-e2e-00000009",
                "executorId": executor_id,
                "environmentSerials": ["4902"],
            },
            headers=CSRF,
        )
        assert scan_created.status_code == 202, scan_created.text
        scan_task_id = scan_created.json()["data"]["taskId"]
        _, scan_token = _lease_and_start(
            device_client, credential, expect_type="after.sale.scan.v1")
        scan_finish = device_client.post(
            f"/v1/executor-channel/tasks/{scan_task_id}/finish",
            json={
                "leaseToken": scan_token,
                "outcome": "succeeded",
                "resultCode": "after_sale_scan_completed",
                "resultSummary": {
                    "rows": [_scan_row("4902", order_no="GSH0001",
                                       claimable=False)],
                    "totalCount": 1,
                    "claimableCount": 0,
                },
            },
            headers=device_headers(credential),
        )
        assert scan_finish.status_code == 200, scan_finish.text
        scan_data = web_client.get(
            f"/v1/after-sale/scan/{scan_task_id}").json()["data"]
        assert scan_data["lastRuns"]["合成店铺-4902"]["refundBillId"] == "RB0001"
        assert scan_data["lastRuns"]["合成店铺-4902"]["runId"] == run_id
        assert scan_data["lastRuns"]["合成店铺-4902"]["status"] == "ok"

        # 截图二进制已随进度落库，且带 7 天过期时间
        with database.session_factory() as session:
            from xynigo_auth.models import AfterSaleClaimResult
            stored = session.scalar(
                select(AfterSaleClaimResult).where(
                    AfterSaleClaimResult.run_id == uuid.UUID(run_id),
                    AfterSaleClaimResult.order_no == "GSH0001",
                )
            )
            assert stored is not None
            assert stored.screenshot_content == JPEG
            assert stored.screenshot_sha256 == sha
            assert stored.screenshot_expires_at is not None
            assert stored.environment_serial == "4902"
            other = session.scalar(
                select(AfterSaleClaimResult).where(
                    AfterSaleClaimResult.run_id == uuid.UUID(run_id),
                    AfterSaleClaimResult.order_no == "GSH0002",
                )
            )
            assert other is not None and other.screenshot_content is None
        break


def test_after_sale_claim_progress_rejects_order_outside_run(tmp_path) -> None:
    """跨批次白名单：orderNo 不在建 Run 清单内必须 422（不落脏行）。"""
    for web_client, device_client, ids, _database in _e2e_setup(tmp_path):
        executor_id, credential = ids["executorId"], ids["credential"]
        heartbeat(device_client, credential, capabilities=AS_CAPABILITIES,
                  client_version=CLIENT_VERSION)
        created = web_client.post(
            "/v1/operation-runs/after-sale-claim",
            json={
                "idempotencyKey": "as-claim-e2e-00000002",
                "executorId": executor_id,
                "items": [{"environmentSerial": "4902", "orderNo": "GSH0001",
                           "storeName": "合成店铺-4902", "packageNo": ""}],
            },
            headers=CSRF,
        )
        assert created.status_code == 202, created.text
        run_id = created.json()["data"]["runId"]
        task_id, lease_token = _lease_and_start(
            device_client, credential, expect_type="after.sale.claim.v1")

        outside = device_client.post(
            f"/v1/executor-channel/tasks/{task_id}/progress",
            json={
                "leaseToken": lease_token,
                "phase": "after_sale.claim.running",
                "current": 1,
                "total": 1,
                "snapshot": {"rows": [_claim_row("GSH9999", "4902")]},
            },
            headers=device_headers(credential),
        )
        assert outside.status_code == 422, outside.text

        mismatch = device_client.post(
            f"/v1/executor-channel/tasks/{task_id}/progress",
            json={
                "leaseToken": lease_token,
                "phase": "after_sale.claim.running",
                "current": 1,
                "total": 1,
                "snapshot": {"rows": [_claim_row("GSH0001", "8888")]},
            },
            headers=device_headers(credential),
        )
        assert mismatch.status_code == 422, mismatch.text

        snapshot = web_client.get(
            f"/v1/operation-runs/after-sale-claim/{run_id}").json()["data"]
        assert snapshot["rows"] == []
        break


# ===== 取消：裸 POST 把 stop 请求下发到任务 =====
def test_after_sale_scan_cancel_requests_stop(tmp_path) -> None:
    for web_client, device_client, ids, database in _e2e_setup(tmp_path):
        executor_id, credential = ids["executorId"], ids["credential"]
        heartbeat(device_client, credential, capabilities=AS_CAPABILITIES,
                  client_version=CLIENT_VERSION)
        created = web_client.post(
            "/v1/after-sale/scan",
            json={**_scan_body("4902"), "executorId": executor_id},
            headers=CSRF,
        )
        assert created.status_code == 202, created.text
        task_id = created.json()["data"]["taskId"]
        task_id, lease_token = _lease_and_start(
            device_client, credential, expect_type="after.sale.scan.v1")

        # 前端裸 POST：不带 body、不带 CSRF 以外的任何负载
        cancelled = web_client.post(
            f"/v1/after-sale/scan/{task_id}/cancel", headers=CSRF)
        assert cancelled.status_code == 200, cancelled.text
        data = cancelled.json()["data"]
        assert data["taskId"] == task_id
        assert data["status"] == "cancel_requested"
        assert data["summary"]["totalCount"] == 1

        with database.session_factory() as session:
            from xynigo_auth.models import ExecutorTask
            task = session.get(ExecutorTask, uuid.UUID(task_id))
            assert task is not None
            assert task.cancel_requested_at is not None
            assert task.status == "cancel_requested"

        # 执行器在租约内回终态；已终态任务再次取消：原样返回，不改写结果
        finish = device_client.post(
            f"/v1/executor-channel/tasks/{task_id}/finish",
            json={
                "leaseToken": lease_token,
                "outcome": "succeeded",
                "resultCode": "after_sale_scan_completed",
                "resultSummary": {"rows": [], "totalCount": 1,
                                  "claimableCount": 0},
            },
            headers=device_headers(credential),
        )
        assert finish.status_code == 200, finish.text
        again = web_client.post(
            f"/v1/after-sale/scan/{task_id}/cancel", headers=CSRF)
        assert again.status_code == 200, again.text
        assert again.json()["data"]["status"] == "succeeded"
        break


def test_after_sale_claim_cancel_requests_stop(tmp_path) -> None:
    for web_client, device_client, ids, _database in _e2e_setup(tmp_path):
        executor_id, credential = ids["executorId"], ids["credential"]
        heartbeat(device_client, credential, capabilities=AS_CAPABILITIES,
                  client_version=CLIENT_VERSION)
        created = web_client.post(
            "/v1/operation-runs/after-sale-claim",
            json={
                "idempotencyKey": "as-claim-e2e-00000003",
                "executorId": executor_id,
                "items": [{"environmentSerial": "4902", "orderNo": "GSH0001",
                           "storeName": "合成店铺-4902", "packageNo": ""},
                          {"environmentSerial": "4901", "orderNo": "GSH0002",
                           "storeName": "合成店铺-4901", "packageNo": ""}],
            },
            headers=CSRF,
        )
        assert created.status_code == 202, created.text
        run_id = created.json()["data"]["runId"]
        task_id, _lease_token = _lease_and_start(
            device_client, credential, expect_type="after.sale.claim.v1")

        cancelled = web_client.post(
            f"/v1/operation-runs/after-sale-claim/{run_id}/cancel", headers=CSRF)
        assert cancelled.status_code == 200, cancelled.text
        data = cancelled.json()["data"]
        assert data["stopRequested"] is True
        assert data["status"] == "running"
        assert data["progressTotal"] == 2

        # 未知批次：404
        missing = web_client.post(
            f"/v1/operation-runs/after-sale-claim/{uuid.uuid4()}/cancel",
            headers=CSRF)
        assert missing.status_code == 404, missing.text
        assert web_client.get(
            f"/v1/operation-runs/after-sale-claim/{uuid.uuid4()}"
        ).status_code == 404
        break


# ===== 退款跟踪：建任务 → 进度落跟踪表 → GET 形状 + 白名单 =====
def _track_row(bill, *, order_no="GSH0001", serial="4586", phase="reviewing",
               account="****7935", amount="270.22", status="ok"):
    return {
        "refundBillId": bill, "orderNo": order_no, "environmentSerial": serial,
        "storeName": "合成店铺", "status": status, "phase": phase,
        "phaseLabel": {"reviewing": "审核中", "refunded": "已退款"}.get(phase, phase),
        "countdown": "23:48:25" if phase == "reviewing" else "",
        "refundAccount": account, "amount": amount,
        "checkedAt": "2026-09-16T01:08:17+00:00", "note": None,
        "errorSummary": None, "durationSeconds": 6,
    }


def test_after_sale_track_create_progress_and_whitelist(tmp_path) -> None:
    """④ 退款跟踪：本用例盯的是评审指出的两个真问题——
    路由授权（曾误写 authorize 导致请求 500）与回访行必须落在任务清单内。"""
    for web_client, device_client, ids, database in _e2e_setup(tmp_path):
        executor_id, credential = ids["executorId"], ids["credential"]
        heartbeat(device_client, credential, capabilities=AS_CAPABILITIES,
                  client_version=CLIENT_VERSION)

        created = web_client.post(
            "/v1/after-sale/track",
            json={
                "idempotencyKey": "as-track-e2e-00000001",
                "executorId": executor_id,
                "items": [
                    {"environmentSerial": "4586", "orderNo": "GSH0001",
                     "refundBillId": "2390833880014851"},
                    {"environmentSerial": "4904", "orderNo": "GSH0002",
                     "refundBillId": "2390783198795777"},
                ],
            },
            headers=CSRF,
        )
        assert created.status_code == 202, created.text
        task_id = created.json()["data"]["taskId"]

        leased_id, lease_token = _lease_and_start(
            device_client, credential, expect_type="after.sale.track.v1")
        assert leased_id == task_id

        progress = device_client.post(
            f"/v1/executor-channel/tasks/{task_id}/progress",
            json={
                "leaseToken": lease_token,
                "phase": "after_sale.track.running",
                "current": 2,
                "total": 2,
                "snapshot": {"rows": [
                    _track_row("2390833880014851", order_no="GSH0001"),
                    _track_row("2390783198795777", order_no="GSH0002",
                               serial="4904", phase="refunded",
                               account="****2813", amount="35.26"),
                ]},
            },
            headers=device_headers(credential),
        )
        assert progress.status_code == 200, progress.text

        # GET：形状 + 阶段分布（读的是跟踪表，覆盖更新后的最新态）
        snap = web_client.get(f"/v1/after-sale/track/{task_id}")
        assert snap.status_code == 200, snap.text
        data = snap.json()["data"]
        assert data["taskId"] == task_id
        rows = data["summary"]["rows"]
        assert len(rows) == 2, rows
        assert data["summary"]["counts"] == {"reviewing": 1, "refunded": 1}
        reviewing = next(r for r in rows if r["phase"] == "reviewing")
        assert reviewing["refundAccount"] == "****7935"
        assert reviewing["countdown"] == "23:48:25"

        # 落表：按 refund_bill_id 唯一，值为最近一次回访态
        with database.session_factory() as session:
            from xynigo_auth.models import AfterSaleRefundTracking
            records = session.scalars(select(AfterSaleRefundTracking)).all()
            assert {r.refund_bill_id for r in records} == {
                "2390833880014851", "2390783198795777"}
            one = next(r for r in records
                       if r.refund_bill_id == "2390833880014851")
            assert one.phase == "reviewing"
            assert one.refund_account == "****7935"
            assert one.order_no == "GSH0001"

        # 白名单：清单外的退款单号必须整批 422，不得覆盖别人的行
        outside = device_client.post(
            f"/v1/executor-channel/tasks/{task_id}/progress",
            json={
                "leaseToken": lease_token,
                "phase": "after_sale.track.running",
                "current": 1,
                "total": 1,
                "snapshot": {"rows": [
                    _track_row("9999999999999999", phase="refunded")]},
            },
            headers=device_headers(credential),
        )
        assert outside.status_code == 422, outside.text
        # 终态 finish：回访是 scan 型任务，终态必须能回来（含 phaseCounts 这个 dict）。
        # 曾把 phaseCounts 放进 BUSINESS_RESULT_KEYS 却按非负整数校验 → finish 422，
        # 任务永远 running、④ 停在「回访运行中」。这里锁死该回归。
        finished = device_client.post(
            f"/v1/executor-channel/tasks/{task_id}/finish",
            json={
                "leaseToken": lease_token,
                "outcome": "succeeded",
                "resultCode": "after_sale_track_completed",
                "resultSummary": {
                    "runStatus": "completed",
                    "phase": "after_sale.track.completed",
                    "progressCompleted": 2,
                    "progressTotal": 2,
                    "totalCount": 2,
                    "successCount": 2,
                    "failedCount": 0,
                    "stoppedCount": 0,
                    "phaseCounts": {"reviewing": 1, "refunded": 1},
                    "errorCode": "",
                    "errorSummary": "",
                },
            },
            headers=device_headers(credential),
        )
        assert finished.status_code == 200, finished.text

        terminal = web_client.get(f"/v1/after-sale/track/{task_id}")
        assert terminal.status_code == 200, terminal.text
        assert terminal.json()["data"]["status"] in (
            "succeeded", "completed"), terminal.text
        break


def test_claim_progress_falls_back_to_request_display_fields(tmp_path) -> None:
    """③ 的送达时间 / ④ 的商品图：执行器没上报时，云端用自己清单里的值兜底。

    真机踩过：老执行器的桥接层按固定字段重建条目（0.17/0.18.0 都如此），
    把 deliveredAt/goodsImg 丢掉，于是 ③ 两列恒空、④ 缩略图恒空——而这两个值
    Web 已经发过来了、云端清单里就有。这里断言「执行器一个字都不报」也能落库。
    """
    for web_client, device_client, ids, database in _e2e_setup(tmp_path):
        executor_id, credential = ids["executorId"], ids["credential"]
        heartbeat(device_client, credential, capabilities=AS_CAPABILITIES,
                  client_version=CLIENT_VERSION)

        created = web_client.post(
            "/v1/operation-runs/after-sale-claim",
            json={
                "idempotencyKey": "as-claim-fallback-000001",
                "executorId": executor_id,
                "items": [{
                    "environmentSerial": "4589", "orderNo": "GSH1FALLBACK",
                    "deliveredAt": "04 Sep 2026 16:56:59",
                    "goodsImg": "//img.ltwebstatic.com/v4/j/pi/fallback.jpg",
                }],
            },
            headers=CSRF,
        )
        assert created.status_code == 202, created.text
        run_id = created.json()["data"]["runId"]
        task_id, lease_token = _lease_and_start(
            device_client, credential, expect_type="after.sale.claim.v1")

        # 执行器上报的行里**故意不带**这两个字段（模拟老桥接层）
        progress = device_client.post(
            f"/v1/executor-channel/tasks/{task_id}/progress",
            json={
                "leaseToken": lease_token,
                "phase": "after_sale.claim.running",
                "current": 1, "total": 1,
                "snapshot": {"rows": [{
                    "orderNo": "GSH1FALLBACK", "environmentSerial": "4589",
                    "status": "ok", "refundBillId": "2390000000000001",
                    "refundPath": "Cuenta original de pago",
                    "submittedAt": "2026-09-16T06:16:50+00:00",
                    "note": "", "errorSummary": None,
                    "packageNo": "", "refundAccount": "****0212",
                    "durationSeconds": 6, "screenshotSha256": "",
                }]},
            },
            headers=device_headers(credential),
        )
        assert progress.status_code == 200, progress.text

        with database.session_factory() as session:
            from xynigo_auth.models import AfterSaleClaimResult
            row = session.scalar(select(AfterSaleClaimResult).where(
                AfterSaleClaimResult.run_id == uuid.UUID(run_id)))
            assert row.goods_img == "//img.ltwebstatic.com/v4/j/pi/fallback.jpg"
            assert row.delivered_at == "04 Sep 2026 16:56:59"

        # 快照（③ 的渲染来源）同样带着这两个字段
        snap = web_client.get(
            f"/v1/operation-runs/after-sale-claim/{run_id}").json()["data"]
        assert snap["rows"][0]["goodsImg"] == (
            "//img.ltwebstatic.com/v4/j/pi/fallback.jpg")
        assert snap["rows"][0]["deliveredAt"] == "04 Sep 2026 16:56:59"
        break


# ===== 退款跟踪导出：本任务清单内的行 → xlsx（列序 + 授权） =====
def test_after_sale_track_export_workbook_and_auth(tmp_path) -> None:
    """④ 导出：内容是本次回访任务清单内的行，未登录必须拦在授权层。

    盯两类问题：导出列与工作台 ④ 表头漂移（整列错位自查抓不到），
    以及路由漏授权（第一轮评审在同类路由上抓到过 authorize 写错）。
    """
    for web_client, device_client, ids, database in _e2e_setup(tmp_path):
        executor_id, credential = ids["executorId"], ids["credential"]
        heartbeat(device_client, credential, capabilities=AS_CAPABILITIES,
                  client_version=CLIENT_VERSION)

        created = web_client.post(
            "/v1/after-sale/track",
            json={
                "idempotencyKey": "as-track-export-000001",
                "executorId": executor_id,
                "items": [
                    {"environmentSerial": "4586", "orderNo": "GSH0001",
                     "refundBillId": "2390833880014851"},
                    {"environmentSerial": "4904", "orderNo": "GSH0002",
                     "refundBillId": "2390783198795777"},
                ],
            },
            headers=CSRF,
        )
        assert created.status_code == 202, created.text
        task_id = created.json()["data"]["taskId"]
        leased_id, lease_token = _lease_and_start(
            device_client, credential, expect_type="after.sale.track.v1")
        assert leased_id == task_id

        progress = device_client.post(
            f"/v1/executor-channel/tasks/{task_id}/progress",
            json={
                "leaseToken": lease_token,
                "phase": "after_sale.track.running",
                "current": 2,
                "total": 2,
                "snapshot": {"rows": [
                    _track_row("2390833880014851", order_no="GSH0001",
                               serial="4586", phase="refunded",
                               account="****2813", amount="35.26"),
                    _track_row("2390783198795777", order_no="GSH0002",
                               serial="4904"),
                ]},
            },
            headers=device_headers(credential),
        )
        assert progress.status_code == 200, progress.text

        exported = web_client.get(f"/v1/after-sale/track/{task_id}/export")
        assert exported.status_code == 200, exported.text
        assert exported.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument"
            ".spreadsheetml.sheet")
        assert quote("退款跟踪结果_") in exported.headers["content-disposition"]
        assert exported.headers["x-xynigo-row-count"] == "2"

        sheet = load_workbook(BytesIO(exported.content)).active
        assert [cell.value for cell in sheet[1]] == list(EXPORT_HEADERS)
        assert sheet.max_row == 3
        # 逐行成对校验：订单号与退款单号必须同属一行（错位就是把两列拆开了）
        assert {
            sheet.cell(row=index, column=2).value:
            sheet.cell(row=index, column=4).value for index in (2, 3)
        } == {"GSH0001": "2390833880014851",
              "GSH0002": "2390783198795777"}
        assert {
            sheet.cell(row=index, column=7).value for index in (2, 3)
        } == {"审核中", "历史退款状态 · 待回访"}
        # 商品图来自提交结果表：本用例没提交过，保持空列
        assert [sheet.cell(row=index, column=3).value
                for index in (2, 3)] == [None, None]

        # 别的任务导不出来：不在清单里的任务号一律 404
        assert web_client.get(
            f"/v1/after-sale/track/{uuid.uuid4()}/export"
        ).status_code == 404

        # 未登录必须 401
        web_client.cookies.clear()
        assert web_client.get(
            f"/v1/after-sale/track/{task_id}/export"
        ).status_code == 401
        break


# ===== 提交历史：列表 / 详情 / 导出 / 可见范围 =====
def test_after_sale_claim_history_scope_paging_detail_export(tmp_path) -> None:
    """提交历史：**租户内互相可见**（含他人批次）、cursor 分页、批次导出、重提来源。

    可见范围是刻意偏离物流/建环境的「本人 + 管理员」口径（需求 §6.3），
    所以这里必须有「别人的批次也看得到」这条断言，否则等于没验收。
    """
    for web_client, device_client, ids, database in _e2e_setup(tmp_path):
        executor_id, credential = ids["executorId"], ids["credential"]
        heartbeat(device_client, credential, capabilities=AS_CAPABILITIES,
                  client_version=CLIENT_VERSION)

        first = web_client.post(
            "/v1/operation-runs/after-sale-claim",
            json={
                "idempotencyKey": "as-claim-history-0000001",
                "executorId": executor_id,
                "items": [{
                    "environmentSerial": "4589", "orderNo": "GSH1A",
                    "deliveredAt": "04 Sep 2026 16:56:59",
                    "goodsImg": "//img.ltwebstatic.com/v4/j/pi/x.jpg",
                }],
            },
            headers=CSRF,
        )
        assert first.status_code == 202, first.text
        first_run_id = first.json()["data"]["runId"]

        # 同租户、另一个操作人的批次（直接落库，避免再造一套登录）
        with database.session_factory() as session:
            from xynigo_auth.models import AfterSaleClaimRun, User
            tenant_id = session.scalar(
                select(AfterSaleClaimRun.tenant_id).where(
                    AfterSaleClaimRun.id == uuid.UUID(first_run_id))
            )
            peer = User(tenant_id=tenant_id, feishu_open_id="ou_as_peer",
                        display_name="同事甲", status="active")
            session.add(peer)
            session.flush()
            # 把同事的批次设成**更新**：/latest 若按租户取就会取到它，用例才有区分度
            moment = datetime.now(timezone.utc) + timedelta(minutes=5)
            peer_run = AfterSaleClaimRun(
                id=uuid.uuid4(), tenant_id=tenant_id, actor_user_id=peer.id,
                source_run_key="as-claim-history-peer-0001",
                payload_hash="0" * 64, browser_mode="visible",
                status="completed", phase="after_sale.completed",
                progress_completed=1, progress_total=1, total_count=1,
                success_count=1, failed_count=0, skipped_count=0,
                stopped_count=0, request_summary={"items": []},
                source="cloud_web", created_at=moment, updated_at=moment,
            )
            session.add(peer_run)
            session.commit()
            peer_run_id = str(peer_run.id)

        listed = web_client.get(
            "/v1/operation-runs/after-sale-claim/history?limit=20")
        assert listed.status_code == 200, listed.text
        data = listed.json()["data"]
        runs = {item["runId"]: item for item in data["items"]}
        assert first_run_id in runs, runs
        assert peer_run_id in runs, "别人的批次也必须可见（租户内互相可见）"
        assert runs[first_run_id]["totalCount"] == 1
        assert runs[first_run_id]["actorName"] != ""
        assert [actor["displayName"] for actor in data["actors"]] != []

        # ③ 的「恢复最近一批」只恢复本人的：同事的批次更新，也不能被恢复
        latest = web_client.get(
            "/v1/operation-runs/after-sale-claim/latest")
        assert latest.status_code == 200, latest.text
        assert latest.json()["data"]["runId"] == first_run_id, (
            "③ 恢复最近一批必须是本人的（共享视图在历史弹层）")

        # 分页：limit=1 拿第一页 + 游标，第二页不重不漏
        page_one = web_client.get(
            "/v1/operation-runs/after-sale-claim/history?limit=1"
        ).json()["data"]
        assert len(page_one["items"]) == 1 and page_one["hasMore"] is True
        page_two = web_client.get(
            "/v1/operation-runs/after-sale-claim/history?limit=1&cursor="
            + page_one["nextCursor"]
        ).json()["data"]
        second_ids = [item["runId"] for item in page_two["items"]]
        assert page_one["items"][0]["runId"] not in second_ids
        assert set(second_ids) <= set(runs)

        # 状态筛选
        filtered = web_client.get(
            "/v1/operation-runs/after-sale-claim/history?status=completed"
        ).json()["data"]
        assert [item["runId"] for item in filtered["items"]] == [peer_run_id]

        # 详情：表头字段 + 逐单行（含商品图/送达时间，来自 request_summary 与结果表）
        detail = web_client.get(
            f"/v1/operation-runs/after-sale-claim/history/{peer_run_id}")
        assert detail.status_code == 200, detail.text
        batch = detail.json()["data"]["batch"]
        assert batch["runId"] == peer_run_id
        assert batch["actorName"] == "同事甲"
        assert batch["retryFromRunId"] == ""
        assert "rows" in detail.json()["data"]

        # 先把第一批的任务跑完，否则同一执行器会被 executor_task_busy 挡住
        task_id, lease_token = _lease_and_start(
            device_client, credential, expect_type="after.sale.claim.v1")
        finished = device_client.post(
            f"/v1/executor-channel/tasks/{task_id}/finish",
            json={
                "leaseToken": lease_token,
                "outcome": "succeeded",
                "resultCode": "after_sale_completed",
                "resultSummary": {
                    "runStatus": "completed",
                    "phase": "after_sale.completed",
                    "progressCompleted": 1, "progressTotal": 1,
                    "totalCount": 1, "successCount": 1, "failedCount": 0,
                    "skippedCount": 0, "stoppedCount": 0,
                    "errorCode": "", "errorSummary": "",
                },
            },
            headers=device_headers(credential),
        )
        assert finished.status_code == 200, finished.text

        # 重提来源：新批次带上 retryFromRunId，列表与详情都要能读出来
        retry = web_client.post(
            "/v1/operation-runs/after-sale-claim",
            json={
                "idempotencyKey": "as-claim-history-retry-0001",
                "executorId": executor_id,
                "retryFromRunId": first_run_id,
                "items": [{"environmentSerial": "4589", "orderNo": "GSH1B"}],
            },
            headers=CSRF,
        )
        assert retry.status_code == 202, retry.text
        retry_run_id = retry.json()["data"]["runId"]
        retry_item = next(
            item for item in web_client.get(
                "/v1/operation-runs/after-sale-claim/history"
            ).json()["data"]["items"] if item["runId"] == retry_run_id)
        assert retry_item["retryFromRunId"] == first_run_id

        # 批次导出：列与 ③ 一致，行数等于该批次的结果行
        exported = web_client.get(
            f"/v1/operation-runs/after-sale-claim/history/{peer_run_id}/export")
        assert exported.status_code == 200, exported.text
        assert exported.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument"
            ".spreadsheetml.sheet")
        assert quote("售后提交结果_") in exported.headers["content-disposition"]
        sheet = load_workbook(BytesIO(exported.content)).active
        assert [cell.value for cell in sheet[1]] == list(CLAIM_EXPORT_HEADERS)

        # 坏游标要走 422（过期/被改过/别的租户的 id），不能变成未捕获的 500
        bad_cursor = web_client.get(
            "/v1/operation-runs/after-sale-claim/history?cursor="
            + str(uuid.uuid4()))
        assert bad_cursor.status_code == 422, bad_cursor.text
        assert bad_cursor.json()["detail"]["code"] == (
            "after_sale_claim_history_cursor_invalid")

        # 跨租户隔离：别的租户的批次不能在列表里露出来
        with database.session_factory() as session:
            from xynigo_auth.models import (AfterSaleClaimRun, Tenant, User)
            outsider_tenant = Tenant(feishu_tenant_key="tenant_as_history",
                                     name="别的租户", status="active")
            session.add(outsider_tenant)
            session.flush()
            outsider = User(tenant_id=outsider_tenant.id,
                            feishu_open_id="ou_as_outsider",
                            display_name="外租户", status="active")
            session.add(outsider)
            session.flush()
            moment = datetime.now(timezone.utc) - timedelta(minutes=1)
            session.add(AfterSaleClaimRun(
                id=uuid.uuid4(), tenant_id=outsider_tenant.id,
                actor_user_id=outsider.id,
                source_run_key="as-claim-history-outsider-0001",
                payload_hash="1" * 64, browser_mode="visible",
                status="completed", phase="after_sale.completed",
                progress_completed=1, progress_total=1, total_count=1,
                success_count=1, failed_count=0, skipped_count=0,
                stopped_count=0, request_summary={"items": []},
                source="cloud_web", created_at=moment, updated_at=moment,
            ))
            session.commit()
        visible = web_client.get(
            "/v1/operation-runs/after-sale-claim/history?limit=100"
        ).json()["data"]["items"]
        assert all(item["actorName"] != "外租户" for item in visible), (
            "别的租户的批次不得可见")

        # retryFromRunId 只用于展示，不参与幂等：同幂等键改这个字段仍算同一请求
        again = web_client.post(
            "/v1/operation-runs/after-sale-claim",
            json={
                "idempotencyKey": "as-claim-history-retry-0001",
                "executorId": executor_id,
                "retryFromRunId": "",
                "items": [{"environmentSerial": "4589", "orderNo": "GSH1B"}],
            },
            headers=CSRF,
        )
        assert again.status_code == 202, again.text
        assert again.json()["data"]["runId"] == retry_run_id
        # 幂等命中与否看审计的 unchanged（响应体是批次快照，不带这个字段）
        with database.session_factory() as session:
            from xynigo_auth.models import AuditEvent
            replay = session.scalar(
                select(AuditEvent.change_summary).where(
                    AuditEvent.action == "assistant.after_sale.claim.create",
                    AuditEvent.business_object_id == retry_run_id,
                ).order_by(AuditEvent.created_at.desc()).limit(1))
        assert replay.get("unchanged") is True, (
            "重提来源不该参与幂等（补填/改填来源不该 409）")

        # 未登录必须拦在授权层
        web_client.cookies.clear()
        assert web_client.get(
            "/v1/operation-runs/after-sale-claim/history").status_code == 401
        break


def test_after_sale_claim_history_environment_count(tmp_path) -> None:
    """批次「环境数」＝结果行去重 environment_serial（不是请求里的环境数）。

    同一环境的多个订单只能算一个环境；还没有结果行时是 0，
    与请求订单数独立统计。覆盖逐步回传时的 0、1、2 三种取值。
    """
    for web_client, device_client, ids, database in _e2e_setup(tmp_path):
        executor_id, credential = ids["executorId"], ids["credential"]
        heartbeat(device_client, credential, capabilities=AS_CAPABILITIES,
                  client_version=CLIENT_VERSION)

        created = web_client.post(
            "/v1/operation-runs/after-sale-claim",
            json={
                "idempotencyKey": "as-claim-envcount-000001",
                "executorId": executor_id,
                "items": [
                    {"environmentSerial": "4589", "orderNo": "GSH1A"},
                    {"environmentSerial": "4589", "orderNo": "GSH1B"},
                    {"environmentSerial": "4590", "orderNo": "GSH1C"},
                ],
            },
            headers=CSRF,
        )
        assert created.status_code == 202, created.text
        run_id = created.json()["data"]["runId"]

        def history_item():
            items = web_client.get(
                "/v1/operation-runs/after-sale-claim/history"
            ).json()["data"]["items"]
            return next(item for item in items if item["runId"] == run_id)

        # 还没跑：请求里有 2 个环境 3 单，但结果行数为 0 → 环境数 0
        assert history_item()["environmentCount"] == 0
        assert history_item()["totalCount"] == 3

        task_id, lease_token = _lease_and_start(
            device_client, credential, expect_type="after.sale.claim.v1")
        partial = device_client.post(
            f"/v1/executor-channel/tasks/{task_id}/progress",
            json={
                "leaseToken": lease_token,
                "phase": "after_sale.claim.running",
                "current": 2, "total": 3,
                "snapshot": {"rows": [
                    _claim_row("GSH1A", "4589"),
                    _claim_row("GSH1B", "4589"),
                ]},
            },
            headers=device_headers(credential),
        )
        assert partial.status_code == 200, partial.text
        assert history_item()["environmentCount"] == 1
        assert history_item()["totalCount"] == 3
        progress = device_client.post(
            f"/v1/executor-channel/tasks/{task_id}/progress",
            json={
                "leaseToken": lease_token,
                "phase": "after_sale.claim.running",
                "current": 3,
                "total": 3,
                "snapshot": {"rows": [
                    _claim_row("GSH1A", "4589"),
                    _claim_row("GSH1B", "4589"),
                    _claim_row("GSH1C", "4590"),
                ]},
            },
            headers=device_headers(credential),
        )
        assert progress.status_code == 200, progress.text

        item = history_item()
        assert item["environmentCount"] == 2, (
            "同一环境的两个订单只能算一个环境："
            f"{item['environmentCount']} != 2")
        assert item["totalCount"] == 3, "单数照旧按行计，不跟着环境数走"
        # 详情 API 的 batch.environmentCount 与列表同一口径
        detail = web_client.get(
            f"/v1/operation-runs/after-sale-claim/history/{run_id}")
        assert detail.status_code == 200, detail.text
        assert detail.json()["data"]["batch"]["environmentCount"] == 2
        # ③ 的实时快照不吃这个字段（它是批次表头字段，不是提交行字段）
        live = web_client.get(
            f"/v1/operation-runs/after-sale-claim/{run_id}").json()["data"]
        assert "environmentCount" not in live
        break


# ===== 提交：建 Run（幂等）→ 进度 + 截图 → 终态快照 =====


# ===== 按环境单遍直提：动态发现订单 + 环境级结果落库 =====
def _env_row(serial: str, *, status: str, submitted: int = 0,
             blocked: int = 0, failed: int = 0, entry: int = 0,
             note: str = "") -> dict[str, object]:
    """执行器投影后的环境行（闭集，字段与 AfterSaleClaimEnvironmentRow 对齐）。"""
    return {
        "environmentSerial": serial,
        "environmentId": "C" + serial,
        "storeName": "合成店铺-" + serial,
        "accountName": "buyer@example.test",
        "status": status,
        "entryCount": entry,
        "submittedCount": submitted,
        "blockedCount": blocked,
        "failedCount": failed,
        "note": note,
        "errorSummary": None,
        "durationSeconds": 55,
    }


def test_after_sale_claim_environments_direct_write_and_env_results(tmp_path) -> None:
    for web_client, device_client, ids, database in _e2e_setup(tmp_path):
        executor_id, credential = ids["executorId"], ids["credential"]
        heartbeat(device_client, credential, capabilities=AS_CAPABILITIES,
                  client_version=CLIENT_VERSION)
        serials = ["4589", "4590", "4591", "4592"]
        body = {
            "idempotencyKey": "as-claim-env-e2e-000001",
            "executorId": executor_id,
            "browserMode": "visible",
            "environmentSerials": serials,
        }
        created = web_client.post(
            "/v1/operation-runs/after-sale-claim", json=body, headers=CSRF)
        assert created.status_code == 202, created.text
        snapshot = created.json()["data"]
        run_id = snapshot["runId"]
        # 进度以环境计；订单事先未知，因此不挂计划清单
        assert snapshot["totalCount"] == 4
        assert snapshot["progressTotal"] == 4
        assert snapshot["submitMode"] == "environments"
        assert snapshot["environments"] == []
        assert snapshot["rows"] == []

        replay = web_client.post(
            "/v1/operation-runs/after-sale-claim", json=body, headers=CSRF)
        assert replay.status_code == 202, replay.text
        assert replay.json()["data"]["runId"] == run_id

        task_id, lease_token = _lease_and_start(
            device_client, credential, expect_type="after.sale.claim.v1")

        # 订单号执行中动态发现：不在任何请求清单里，也必须能落库
        progress = device_client.post(
            f"/v1/executor-channel/tasks/{task_id}/progress",
            json={
                "leaseToken": lease_token,
                "phase": "after_sale.claim.running",
                "current": 2,
                "total": 4,
                "snapshot": {
                    "rows": [
                        _claim_row("GSH1A", "4589"),
                        _claim_row("GSH1B", "4589"),
                        _claim_row("GSH1C", "4590"),
                        _claim_row("GSH1D", "4590", status="blocked"),
                    ],
                    "environments": [
                        _env_row("4589", status="ok", submitted=2, entry=2),
                        _env_row("4590", status="ok", submitted=1, blocked=1,
                                 entry=2),
                        _env_row("4591", status="running", entry=0),
                        _env_row("4592", status="queued", entry=0),
                    ],
                },
            },
            headers=device_headers(credential),
        )
        assert progress.status_code == 200, progress.text

        partial = web_client.get(
            f"/v1/operation-runs/after-sale-claim/{run_id}").json()["data"]
        assert partial["submitMode"] == "environments"
        assert len(partial["rows"]) == 4
        assert {row["environmentSerial"] for row in partial["rows"]} == {
            "4589", "4590"}
        envs = {row["environmentSerial"]: row for row in partial["environments"]}
        assert envs["4589"]["status"] == "ok"
        assert envs["4589"]["submittedCount"] == 2
        assert envs["4591"]["status"] == "running"

        # 环境行必须在计划清单内：串批次的环境不许混进来
        out_of_scope = device_client.post(
            f"/v1/executor-channel/tasks/{task_id}/progress",
            json={
                "leaseToken": lease_token,
                "phase": "after_sale.claim.running",
                "current": 2,
                "total": 4,
                "snapshot": {"rows": [
                    {**_claim_row("GSH1Z", "9999")}]},
            },
            headers=device_headers(credential),
        )
        assert out_of_scope.status_code == 422, out_of_scope.text
        # 422 不能落脏行：环境清单外的订单号仍不得出现
        after_reject = web_client.get(
            f"/v1/operation-runs/after-sale-claim/{run_id}").json()["data"]
        assert "GSH1Z" not in {row["orderNo"] for row in after_reject["rows"]}

        # 终态：订单数多于环境数也不能被截断（success 5 > total 4）
        finished = device_client.post(
            f"/v1/executor-channel/tasks/{task_id}/finish",
            json={
                "leaseToken": lease_token,
                "outcome": "succeeded",
                "resultCode": "after_sale_completed",
                "resultSummary": {
                    "runStatus": "completed",
                    "phase": "after_sale.completed",
                    "totalCount": 4,
                    "progressTotal": 4,
                    "progressCompleted": 4,
                    "successCount": 5,
                    "failedCount": 0,
                    "skippedCount": 1,
                    "stoppedCount": 0,
                    "uncertainCount": 0,
                    "rows": [
                        _claim_row("GSH1A", "4589"),
                        _claim_row("GSH1B", "4589"),
                        _claim_row("GSH1C", "4590"),
                        _claim_row("GSH1D", "4590", status="blocked"),
                        _claim_row("GSH1E", "4590"),
                        _claim_row("GSH1F", "4590"),
                    ],
                    "environments": [
                        _env_row("4589", status="ok", submitted=2, entry=2),
                        _env_row("4590", status="ok", submitted=3, blocked=1,
                                 entry=4),
                        _env_row("4591", status="skip", entry=0,
                                 note="未发现丢件退款入口，环境跳过"),
                        _env_row("4592", status="skip", entry=0,
                                 note="未发现丢件退款入口，环境跳过"),
                    ],
                },
            },
            headers=device_headers(credential),
        )
        assert finished.status_code == 200, finished.text

        final = web_client.get(
            f"/v1/operation-runs/after-sale-claim/{run_id}").json()["data"]
        assert final["status"] == "completed"
        assert final["successCount"] == 5, "订单级计数不能被环境数截断"
        assert final["skippedCount"] == 1
        envs = {row["environmentSerial"]: row for row in final["environments"]}
        assert envs["4591"]["status"] == "skip"
        assert "入口" in envs["4591"]["note"]
        assert len(envs) == 4

        # 历史列表/详情：submitMode 透出，环境级结果随详情返回
        history = web_client.get(
            "/v1/operation-runs/after-sale-claim/history").json()["data"]
        item = next(entry for entry in history["items"]
                    if entry["runId"] == run_id)
        assert item["submitMode"] == "environments"
        detail = web_client.get(
            f"/v1/operation-runs/after-sale-claim/history/{run_id}").json()["data"]
        assert detail["batch"]["submitMode"] == "environments"
        assert len(detail["environments"]) == 4
        break


def test_after_sale_claim_environments_requires_capability_and_valid_scope(tmp_path) -> None:
    for web_client, device_client, ids, _database in _e2e_setup(tmp_path):
        executor_id, credential = ids["executorId"], ids["credential"]
        base = {
            "idempotencyKey": "as-claim-env-e2e-000002",
            "executorId": executor_id,
            "browserMode": "visible",
        }
        # 两种范围二选一：都给 / 都不给都必须被契约拦下
        both = web_client.post("/v1/operation-runs/after-sale-claim", headers=CSRF,
                              json={**base, "environmentSerials": ["4589"],
                                    "items": [{"environmentSerial": "4589",
                                               "orderNo": "GSH1A"}]})
        assert both.status_code == 422, both.text
        neither = web_client.post("/v1/operation-runs/after-sale-claim",
                                 headers=CSRF, json=dict(base))
        assert neither.status_code == 422, neither.text
        duplicated = web_client.post(
            "/v1/operation-runs/after-sale-claim", headers=CSRF,
            json={**base, "environmentSerials": ["4589", "4589"]})
        assert duplicated.status_code == 422, duplicated.text

        # 老执行器（无 claim-environment 能力位）不能接直提任务
        legacy = [cap for cap in AS_CAPABILITIES
                  if cap != "after.sale.claim-environment.v1"]
        heartbeat(device_client, credential, capabilities=legacy,
                  client_version=CLIENT_VERSION)
        denied = web_client.post(
            "/v1/operation-runs/after-sale-claim", headers=CSRF,
            json={**base, "environmentSerials": ["4589"]})
        assert denied.status_code == 409, denied.text
        assert "executor_after_sale_environment_upgrade_required" in denied.text
        break


def test_env_results_migration_defaults_existing_rows():
    """0044 迁移：老批次补空列表默认值，不破坏既有行。"""
    import runpy
    from pathlib import Path
    import sqlalchemy as sa
    from alembic.migration import MigrationContext
    from alembic.operations import Operations
    migration = runpy.run_path(str(Path(__file__).resolve().parents[1]
                                   / 'migrations/versions'
                                   / '0044_after_sale_env_results.py'))
    assert len(migration['revision']) <= 32
    engine = sa.create_engine('sqlite://')
    with engine.begin() as connection:
        connection.execute(sa.text(
            'CREATE TABLE after_sale_claim_runs (id INTEGER PRIMARY KEY)'))
        connection.execute(sa.text(
            'INSERT INTO after_sale_claim_runs (id) VALUES (1)'))
        with Operations.context(MigrationContext.configure(connection)):
            migration['upgrade']()
        value = connection.execute(sa.text(
            'SELECT environment_results FROM after_sale_claim_runs'
            ' WHERE id=1')).scalar()
        assert value == '[]'
        with Operations.context(MigrationContext.configure(connection)):
            migration['downgrade']()
