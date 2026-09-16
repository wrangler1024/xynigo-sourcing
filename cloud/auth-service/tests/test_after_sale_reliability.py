"""Atomic final snapshots and cancellation, using isolated synthetic database."""
import io
from openpyxl import load_workbook
from sqlalchemy import select
from xynigo_auth.models import AfterSaleRefundTracking
from test_after_sale_claim import _e2e_setup, _lease_and_start, _track_row, AS_CAPABILITIES, CLIENT_VERSION
from test_executor_channel import CSRF, device_headers, heartbeat


def test_final_receipts_packages_failed_read_and_track_cancel(tmp_path):
    for web, device, ids, database in _e2e_setup(tmp_path):
        credential = ids['credential']
        heartbeat(device, credential, capabilities=AS_CAPABILITIES, client_version=CLIENT_VERSION)
        created = web.post('/v1/operation-runs/after-sale-claim', headers=CSRF, json={
            'executorId':ids['executorId'], 'idempotencyKey':'synthetic-reliable-claim',
            'items':[{'environmentSerial':'ENV','orderNo':'ORDER'}]})
        assert created.status_code == 202, created.text
        run = created.json()['data']['runId']
        task, token = _lease_and_start(device, credential, expect_type='after.sale.claim.v1')
        row = {'environmentSerial':'ENV','orderNo':'ORDER','status':'fail',
               'refundBillId':'B2','refunds':[
                   {'packageNo':'P1','refundBillId':'B1'}, {'packageNo':'P2','refundBillId':'B2'}],
               'errorSummary':'remaining package failed'}
        body = {'leaseToken':token,'outcome':'failed','resultCode':'after_sale_failed',
                'resultSummary':{'runStatus':'failed','phase':'after_sale.failed',
                    'totalCount':1,'progressTotal':1,'progressCompleted':1,
                    'successCount':0,'failedCount':1,'rows':[row]}}
        invalid = {**body, 'resultSummary':{**body['resultSummary'], 'rows':[{**row,'orderNo':'OTHER'}]}}
        rejected = device.post(f'/v1/executor-channel/tasks/{task}/finish', headers=device_headers(credential),json=invalid)
        assert rejected.status_code == 422
        assert web.get(f'/v1/operation-runs/after-sale-claim/{run}').json()['data']['status'] == 'running'
        for _ in range(2):
            done = device.post(f'/v1/executor-channel/tasks/{task}/finish', headers=device_headers(credential),json=body)
            assert done.status_code == 200, done.text
        for path in [f'/v1/operation-runs/after-sale-claim/{run}',f'/v1/operation-runs/after-sale-claim/history/{run}']:
            snapshot = web.get(path).json()['data']
            assert [r['refundBillId'] for r in snapshot['rows'][0]['refunds']] == ['B1','B2']
        export = web.get(f'/v1/operation-runs/after-sale-claim/history/{run}/export')
        assert export.status_code == 200
        assert load_workbook(io.BytesIO(export.content)).active.cell(2,6).value == 'B1\nB2'
        items = [{'environmentSerial':'ENV','orderNo':'ORDER','refundBillId':'B1'}]
        def create_track(key):
            r = web.post('/v1/after-sale/track', headers=CSRF, json={
                'executorId':ids['executorId'],'idempotencyKey':key,'items':items})
            assert r.status_code == 202, r.text
            return r.json()['data']['taskId']
        track = create_track('synthetic-reliable-track')
        _, lease = _lease_and_start(device, credential, expect_type='after.sale.track.v1')
        good = _track_row('B1',order_no='ORDER',serial='ENV')
        progress = device.post(f'/v1/executor-channel/tasks/{track}/progress',headers=device_headers(credential),json={
            'leaseToken':lease,'phase':'after_sale.track.running','current':1,'total':1,'snapshot':{'rows':[good]}})
        assert progress.status_code == 200, progress.text
        done = device.post(f'/v1/executor-channel/tasks/{track}/finish',headers=device_headers(credential),json={
            'leaseToken':lease,'outcome':'failed','resultCode':'after_sale_track_failed',
            'resultSummary':{'runStatus':'failed','phase':'after_sale.track.failed','totalCount':1,
                'progressTotal':1,'progressCompleted':1,'successCount':0,'failedCount':1,
                'rows':[{**items[0],'status':'fail','errorSummary':'unrecognized page'}]}})
        assert done.status_code == 200, done.text
        assert web.get(f'/v1/after-sale/track/{track}').json()['data']['summary']['rows'][0]['status']=='fail'
        with database.session_factory() as session:
            record = session.scalar(select(AfterSaleRefundTracking).where(AfterSaleRefundTracking.refund_bill_id=='B1'))
            assert record.phase == good['phase'] and record.amount == good['amount']
        cancelled = create_track('synthetic-reliable-cancel')
        assert web.post(f'/v1/after-sale/track/{task}/cancel',headers=CSRF).status_code == 404
        result = web.post(f'/v1/after-sale/track/{cancelled}/cancel',headers=CSRF)
        assert result.status_code == 200, result.text
        assert result.json()['data']['status']=='cancelled'
        old = [cap for cap in AS_CAPABILITIES if cap != 'after.sale.reliable-results.v1']
        heartbeat(device, credential, capabilities=old, client_version=CLIENT_VERSION)
        denied = web.post('/v1/after-sale/track', headers=CSRF, json={
            'executorId':ids['executorId'],'idempotencyKey':'synthetic-legacy-denied','items':items})
        assert denied.status_code == 409
        assert 'executor_after_sale_reliability_upgrade_required' in denied.text
        break


def test_refund_migration_preserves_existing_rows_and_defaults():
    import runpy
    from pathlib import Path
    import sqlalchemy as sa
    from alembic.migration import MigrationContext
    from alembic.operations import Operations
    migration = runpy.run_path(str(Path(__file__).resolve().parents[1]/'migrations/versions/0040_after_sale_refunds.py'))
    engine = sa.create_engine('sqlite://')
    with engine.begin() as connection:
        connection.execute(sa.text('CREATE TABLE after_sale_claim_results (id INTEGER PRIMARY KEY)'))
        connection.execute(sa.text('INSERT INTO after_sale_claim_results (id) VALUES (1)'))
        with Operations.context(MigrationContext.configure(connection)):
            migration['upgrade']()
        assert connection.execute(sa.text('SELECT refunds FROM after_sale_claim_results WHERE id=1')).scalar() == '[]'
        connection.execute(sa.text('INSERT INTO after_sale_claim_results (id) VALUES (2)'))
        assert connection.execute(sa.text('SELECT refunds FROM after_sale_claim_results WHERE id=2')).scalar() == '[]'
        with Operations.context(MigrationContext.configure(connection)):
            migration['downgrade']()
        assert connection.execute(sa.text('SELECT COUNT(*) FROM after_sale_claim_results')).scalar() == 2
