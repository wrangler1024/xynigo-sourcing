import io
import runpy
from pathlib import Path

import pytest
import sqlalchemy as sa
from alembic.migration import MigrationContext
from alembic.operations import Operations
from openpyxl import load_workbook

from test_after_sale_claim import _e2e_setup, _lease_and_start, AS_CAPABILITIES, CLIENT_VERSION
from test_executor_channel import CSRF, device_headers, heartbeat


def test_uncertain_receipts_survive_final_history_export_and_block_direct_replay(tmp_path):
    for web, device, ids, database in _e2e_setup(tmp_path):
        credential = ids['credential']
        heartbeat(device, credential, capabilities=AS_CAPABILITIES, client_version=CLIENT_VERSION)
        request = {'executorId':ids['executorId'], 'idempotencyKey':'synthetic-receipt',
                   'items':[{'environmentSerial':'ENV','orderNo':'ORDER'}]}
        created = web.post('/v1/operation-runs/after-sale-claim', headers=CSRF, json=request)
        assert created.status_code == 202, created.text
        run = created.json()['data']['runId']
        task, token = _lease_and_start(device, credential, expect_type='after.sale.claim.v1')
        row = {'environmentSerial':'ENV','orderNo':'ORDER','status':'uncertain',
            'errorSummary':'提交结果待核对', 'refundBillId':'12345', 'refunds':[
                {'refundBillId':'12345','source':'recovered','phaseLabel':'审核中',
                 'applicationTimeText':'7 Sep 2026 12:00:00','timeZone':'America/Mexico_City'}]}
        body = {'leaseToken':token,'outcome':'failed','resultCode':'after_sale_uncertain',
            'resultSummary':{'runStatus':'uncertain','phase':'after_sale.uncertain','totalCount':1,
                'progressTotal':1,'progressCompleted':1,'successCount':0,'failedCount':0,'uncertainCount':1,'rows':[row]}}
        for _ in range(2):
            done = device.post(f'/v1/executor-channel/tasks/{task}/finish', headers=device_headers(credential), json=body)
            assert done.status_code == 200, done.text
        for path in [f'/v1/operation-runs/after-sale-claim/{run}', f'/v1/operation-runs/after-sale-claim/history/{run}']:
            snapshot = web.get(path).json()['data']
            assert snapshot['uncertainCount'] == 1 and snapshot['failedCount'] == 0
            record = snapshot['rows'][0]['refunds'][0]
            assert record['source'] == 'recovered' and record.get('submittedAt', '') == ''
            assert record['applicationTimeText'] == '7 Sep 2026 12:00:00'
        export = web.get(f'/v1/operation-runs/after-sale-claim/history/{run}/export')
        sheet = load_workbook(io.BytesIO(export.content)).active
        assert sheet.cell(2,9).value == '待核对 · 请勿补提 · 平台：审核中'
        assert 'America/Mexico_City' in sheet.cell(2,11).value
        # Same idempotency key still reads the original run, without creating writes.
        same = web.post('/v1/operation-runs/after-sale-claim', headers=CSRF, json=request)
        assert same.status_code in (200,202), same.text
        replay = web.post('/v1/operation-runs/after-sale-claim', headers=CSRF,
                          json={**request,'idempotencyKey':'new-direct-input'})
        assert replay.status_code == 409 and 'after_sale_receipt_reconciliation_required' in replay.text
        alias = web.post('/v1/operation-runs/after-sale-claim', headers=CSRF,
            json={**request,'idempotencyKey':'new-environment-alias','items':[{'environmentSerial':'ALIAS','orderNo':'ORDER'}]})
        assert alias.status_code == 409 and 'after_sale_receipt_reconciliation_required' in alias.text
        canonical = web.post('/v1/operation-runs/after-sale-claim', headers=CSRF,
            json={**request,'idempotencyKey':'new-order-case','items':[{'environmentSerial':'ALIAS','orderNo':' order '}]})
        assert canonical.status_code == 409 and 'after_sale_receipt_reconciliation_required' in canonical.text
        old_caps = [cap for cap in AS_CAPABILITIES if cap != 'after.sale.receipt-recovery.v1']
        heartbeat(device, credential, capabilities=old_caps, client_version=CLIENT_VERSION)
        old = web.post('/v1/operation-runs/after-sale-claim', headers=CSRF,
            json={**request,'idempotencyKey':'old-device-new-order','items':[{'environmentSerial':'ENV','orderNo':'NEW'}]})
        assert old.status_code == 409 and 'executor_after_sale_receipt_upgrade_required' in old.text
        break


def test_status_migration_preserves_history_and_refuses_unsafe_rollback():
    migration = runpy.run_path(str(Path(__file__).resolve().parents[1]/'migrations/versions/0041_after_sale_uncertain.py'))
    engine = sa.create_engine('sqlite://')
    with engine.begin() as connection:
        connection.execute(sa.text('CREATE TABLE after_sale_claim_results (id INTEGER PRIMARY KEY, status TEXT NOT NULL, CONSTRAINT ck_after_sale_result_status CHECK (' + migration['OLD'] + '))'))
        connection.execute(sa.text("INSERT INTO after_sale_claim_results VALUES (1, 'blocked')"))
        with Operations.context(MigrationContext.configure(connection)):
            migration['upgrade']()
        connection.execute(sa.text("INSERT INTO after_sale_claim_results VALUES (2, 'uncertain')"))
        assert connection.execute(sa.text('SELECT COUNT(*) FROM after_sale_claim_results')).scalar() == 2
        with Operations.context(MigrationContext.configure(connection)):
            with pytest.raises(RuntimeError, match='Resolve uncertain'):
                migration['downgrade']()
        connection.execute(sa.text("UPDATE after_sale_claim_results SET status='blocked' WHERE id=2"))
        with Operations.context(MigrationContext.configure(connection)):
            migration['downgrade']()
        assert connection.execute(sa.text('SELECT COUNT(*) FROM after_sale_claim_results')).scalar() == 2


@pytest.mark.parametrize('value', [-1, True, '1'])
def test_uncertain_count_rejects_non_integer_or_negative_values(value):
    from types import SimpleNamespace
    from xynigo_auth.executor_contract import ExecutorTaskFinishBody
    from xynigo_auth.executor_service import ExecutorChannelService, ExecutorServiceError
    body = ExecutorTaskFinishBody(leaseToken='synthetic-' * 4, outcome='failed',
        resultCode='after_sale_uncertain', resultSummary={'runStatus':'uncertain',
            'phase':'after_sale.uncertain','uncertainCount':value})
    with pytest.raises(ExecutorServiceError):
        ExecutorChannelService._validate_business_result(SimpleNamespace(task_type='after.sale.claim.v1'),body)
