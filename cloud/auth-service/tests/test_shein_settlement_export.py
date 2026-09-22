# -*- coding: utf-8 -*-
"""结算看板 xlsx 导出：内容、数值格式、缺汇率提示（全部合成数据）。

导出是财务真正拿去用的东西，所以这里钉的不只是"能生成文件"，而是：
- **金额必须是数值单元格**（带 #,##0.00 格式）——财务要能直接求和、透视，
  写成字符串会让 Excel 当成文本；
- **失败行金额留空**，不参与汇总，也不能显示成 0；
- **缺汇率时人民币合计为空并在页脚点名币种**——不能悄悄少算一个币种。
"""
from __future__ import annotations

import io
from datetime import date, datetime, timedelta
from decimal import Decimal

from openpyxl import load_workbook

from xynigo_auth.shein_settlement_export import build_settlement_export
from xynigo_auth.shein_settlement_service import (
    PayoutBatch,
    StoreSettlementInput,
    build_settlement_summary,
)
from xynigo_auth.shein_settlement_windows import PLATFORM_TZ

NOW = datetime(2026, 9, 17, 12, 0, tzinfo=PLATFORM_TZ)
FX = {"MXN": Decimal("0.3915")}


# 「下次结算」按真实时钟动态判定：导出/汇总在这里自取今天,打款日
# 必须相对运行时构造未来日期,写死日历日次日起变红（20260922 复现）。
D_PAY1 = date.today() + timedelta(days=3)
D_PAY2 = date.today() + timedelta(days=10)
S_PAY1, S_PAY2 = D_PAY1.isoformat(), D_PAY2.isoformat()


def _sheet(summary):
    content, filename, mime = build_settlement_export(summary)
    assert mime.startswith("application/vnd.openxmlformats")
    assert filename.endswith(".xlsx")
    return load_workbook(io.BytesIO(content)).active, content, filename


def _store(name, currency, **kwargs):
    return StoreSettlementInput(
        store_id=name, store_name=name, mode=kwargs.pop("mode", "self"),
        currency=currency, status=kwargs.pop("status", "ok"),
        error_summary=kwargs.pop("error_summary", ""),
        in_transit_amount=kwargs.pop("in_transit", None),
        settled_cumulative_amount=kwargs.pop("settled", None),
        payout_batches=kwargs.pop("batches", ()),
        payment_method=None, synced_at=kwargs.pop("synced_at", NOW),
    )


def _summary(stores, fx=None):
    return build_settlement_summary(stores, FX if fx is None else fx)


def test_export_money_cells_are_numbers_not_text():
    """财务要能在 Excel 里直接求和；字符串会让单元格变文本。"""
    sheet, _, _ = _sheet(_summary([
        _store("甲店", "MXN", in_transit=Decimal("1234.56"),
               settled=Decimal("2000.00"),
               batches=(PayoutBatch(D_PAY1, "MXN", Decimal("500.25")),)),
    ]))
    detail = [c for c in sheet[2]]
    assert isinstance(detail[3].value, (int, float))
    assert detail[3].value == 1234.56
    assert detail[3].number_format == "#,##0.00"
    assert detail[5].value == 500.25
    assert detail[7].value == 2000.0


def test_failed_store_row_has_blank_amounts_and_reason():
    """失败行金额留空（不是 0），并带上失败原因。"""
    sheet, _, _ = _sheet(_summary([
        _store("故障店", "MXN", status="fail",
               error_summary="接口错误：签名错误:生成的签名不正确"),
    ]))
    row = sheet[2]
    assert row[3].value is None
    assert row[4].value is None
    assert row[9].value == "同步失败"
    assert "签名错误" in row[10].value


def test_summary_block_sums_per_currency_and_cny_row():
    sheet, _, _ = _sheet(_summary([
        _store("甲店", "MXN", settled=Decimal("100.00")),
        _store("乙店", "MXN", settled=Decimal("50.00")),
    ]))
    rows = [[c.value for c in row] for row in sheet.iter_rows()]
    # 汇总行的特征：店铺/类型两列为空、第 3 列是币种（明细行的前两列有值）
    currency_rows = [r for r in rows
                     if r[0] is None and r[1] is None and r[2] == "MXN"]
    assert currency_rows and currency_rows[0][7] == 150.0
    cny_rows = [r for r in rows if r[2] == "折人民币"]
    assert cny_rows and cny_rows[0][7] == 58.73      # 150 × 0.3915，落到分


def test_missing_rate_leaves_cny_blank_and_names_currency():
    """缺汇率的币种不能按 1:1 计入；页脚要点名是哪个币种。"""
    sheet, _, _ = _sheet(_summary([
        _store("甲店", "BRL", settled=Decimal("100.00")),
    ], fx={}))
    rows = [[c.value for c in row] for row in sheet.iter_rows()]
    cny_rows = [r for r in rows
                if r[0] is None and r[2] == "折人民币"]
    assert cny_rows and cny_rows[0][7] is None
    flat = " ".join(str(c.value) for row in sheet.iter_rows() for c in row
                    if c.value is not None)
    assert "BRL" in flat and "不按 1:1 计入" in flat


def test_schedule_block_lists_each_pay_date_with_native_amounts():
    """结算排期逐批列明——把所有待结算配一个最早打款日是错的。"""
    sheet, _, _ = _sheet(_summary([
        _store("甲店", "MXN", batches=(
            PayoutBatch(D_PAY1, "MXN", Decimal("600.00")),
            PayoutBatch(D_PAY2, "MXN", Decimal("400.00")),
        )),
    ]))
    rows = [[c.value for c in row] for row in sheet.iter_rows()]
    schedule = [r for r in rows if r[2] in (S_PAY1, S_PAY2)
                and r[4] and "MXN:" in str(r[4])]
    assert [r[2] for r in schedule] == [S_PAY1, S_PAY2]
    assert schedule[0][4] == "MXN:600.00"
    assert schedule[0][3] == 234.9                   # 600 × 0.3915


def test_pay_dates_column_lists_all_batches():
    """明细的预计打款日列出该店全部批次，不只最近一个。"""
    sheet, _, _ = _sheet(_summary([
        _store("甲店", "MXN", batches=(
            PayoutBatch(D_PAY2, "MXN", Decimal("1.00")),
            PayoutBatch(D_PAY1, "MXN", Decimal("2.00")),
        )),
    ]))
    assert sheet[2][6].value == f"{S_PAY1} / {S_PAY2}"


def test_footer_records_data_time_and_coverage():
    sheet, _, _ = _sheet(_summary([
        _store("甲店", "MXN", settled=Decimal("1.00")),
        _store("故障店", "MXN", status="fail", error_summary="x"),
    ]))
    flat = " ".join(str(c.value) for row in sheet.iter_rows() for c in row
                    if c.value is not None)
    assert "数据时间：2026-09-17" in flat
    assert "成功 1 / 共 2 家" in flat


def test_export_with_no_stores_still_produces_valid_file():
    sheet, content, _ = _sheet(_summary([]))
    assert len(content) > 0
    assert sheet[1][0].value == "店铺"          # 表头仍在
