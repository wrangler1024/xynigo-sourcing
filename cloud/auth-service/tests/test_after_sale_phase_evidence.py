"""Five-node facts survive final receipts, failed reads and historical exports."""
import io
from copy import deepcopy

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy import select

from xynigo_auth.models import AfterSaleRefundTracking
from xynigo_auth.operation_contract import AfterSaleTrackRow
from xynigo_auth.operation_service import after_sale_tracking_snapshot
from test_after_sale_claim import _e2e_setup, _lease_and_start, AS_CAPABILITIES, CLIENT_VERSION
from test_executor_channel import CSRF, device_headers, heartbeat


def evidence():
    return {'source':'timeline_nodes_v1','step':2,'title':'Reseña de SHEIN/Vendedor',
            'detail':'Revisión fallida (vendedor/SHEIN)','canSupplement':True,
            'reason':'Comprobantes insuficientes','reasonSource':'negotiation_history'}


def row():
    return {'refundBillId':'12345','orderNo':'ORDER','environmentSerial':'900001',
            'status':'ok','phase':'evidence_required','phaseLabel':'审核未通过 · 待补充凭证',
            'phaseEvidence':evidence(),'checkedAt':'2026-01-01T10:00:00+00:00',
            'note':'平台审核未通过；原因：凭证不足（Comprobantes insuficientes）；可补充凭证重新审核'}


@pytest.mark.parametrize('change',[
    {'phase':'processing'}, {'checkedAt':''}, {'checkedAt':'2026-01-01T10:00:00'},
    {'phaseEvidence':{**evidence(),'canSupplement':False}},
    {'phaseEvidence':{**evidence(),'step':True}},
    {'phaseEvidence':{**evidence(),'reasonSource':''}},
    {'phaseEvidence':{**evidence(),'wholePage':'not allowed'}},
])
def test_contract_rejects_conflicting_and_unbounded_evidence(change):
    with pytest.raises(ValidationError):AfterSaleTrackRow.model_validate({**row(),**change})


def test_old_receipts_remain_accepted_without_relabeling_them_as_bank_complete():
    old={**row(),'phase':'refunded','phaseLabel':'已退款','phaseEvidence':None}
    assert AfterSaleTrackRow.model_validate(old).phase=='refunded'


def test_current_evidence_persists_failed_read_retains_it_and_new_review_clears_old_reason(tmp_path):
    for web,device,ids,database in _e2e_setup(tmp_path):
        credential=ids['credential']
        heartbeat(device,credential,capabilities=AS_CAPABILITIES,client_version=CLIENT_VERSION)
        items=[{k:row()[k] for k in ('refundBillId','orderNo','environmentSerial')}]
        def finish(key, result):
            created=web.post('/v1/after-sale/track',headers=CSRF,json={
                'executorId':ids['executorId'],'idempotencyKey':key,'items':items})
            assert created.status_code==202,created.text
            task,lease=_lease_and_start(device,credential,expect_type='after.sale.track.v1')
            ok=result['status']=='ok'
            response=device.post(f'/v1/executor-channel/tasks/{task}/finish',headers=device_headers(credential),json={
                'leaseToken':lease,'outcome':'succeeded' if ok else 'failed','resultCode':'phase-observation',
                'resultSummary':{'runStatus':'completed' if ok else 'failed','phase':'after_sale.track.completed',
                    'totalCount':1,'progressTotal':1,'progressCompleted':1,'successCount':int(ok),
                    'failedCount':int(not ok),'rows':[result]}})
            assert response.status_code==200,response.text
            return task
        good=finish('synthetic-phase-good',row())
        failed=finish('synthetic-phase-failed',{**items[0],'status':'fail',
                      'errorSummary':'本次状态未确认：页面未完整加载'})
        def read(task):return web.get(f'/v1/after-sale/track/{task}').json()['data']['summary']['rows'][0]
        failed_snapshot=read(failed)
        assert failed_snapshot['status']=='fail'
        assert failed_snapshot['phase']=='evidence_required'
        assert failed_snapshot['phaseEvidence']==evidence()
        assert failed_snapshot['checkedAt'].startswith('2026-01-01T10:00:00')
        assert '凭证不足' in failed_snapshot['note']
        with database.session_factory() as session:
            record=session.scalar(select(AfterSaleRefundTracking))
            snapshot=after_sale_tracking_snapshot(session,record.tenant_id,['12345'])['rows'][0]
            assert snapshot['phaseEvidence']==evidence() and '凭证不足' in snapshot['note']
        next_row=deepcopy(row())
        next_row.update(phase='reviewing',phaseLabel='审核中',note='未终态，下次回访继续跟',checkedAt='2026-01-02T10:00:00+00:00')
        next_row['phaseEvidence'].update(detail='En revisión (vendedor/SHEIN)',canSupplement=False,reason='',reasonSource='')
        reviewed=finish('synthetic-phase-reviewing',next_row)
        assert read(reviewed)['phaseEvidence']['reason']==''
        assert read(failed)==failed_snapshot, 'Later observations must not rewrite an earlier batch'
        export=web.get(f'/v1/after-sale/track/{failed}/export')
        assert export.status_code==200
        sheet=load_workbook(io.BytesIO(export.content)).active
        assert '本次状态未确认' in sheet.cell(2,7).value
        assert '待补充凭证' in sheet.cell(2,7).value
        assert '凭证不足' in sheet.cell(2,10).value
        assert '页面未完整加载' in sheet.cell(2,10).value
        assert read(good)['phaseEvidence']==evidence()


def test_old_executor_cannot_start_new_tracking_with_text_only_phase_detection(tmp_path):
    for web,device,ids,_ in _e2e_setup(tmp_path):
        heartbeat(device,ids['credential'],capabilities=[c for c in AS_CAPABILITIES if c!='after.sale.phase-evidence.v1'],client_version=CLIENT_VERSION)
        response=web.post('/v1/after-sale/track',headers=CSRF,json={
            'executorId':ids['executorId'],'idempotencyKey':'synthetic-old-phase-executor',
            'items':[{k:row()[k] for k in ('refundBillId','orderNo','environmentSerial')}]})
        assert response.status_code==409
        assert 'executor_after_sale_phase_upgrade_required' in response.text


def test_additive_phase_migration_preserves_old_state_and_does_not_fabricate_evidence():
    import runpy
    from pathlib import Path
    import sqlalchemy as sa
    from alembic.migration import MigrationContext
    from alembic.operations import Operations
    migration=runpy.run_path(str(Path(__file__).resolve().parents[1]/'migrations/versions/0043_after_sale_phase_evidence.py'))
    engine=sa.create_engine('sqlite://')
    with engine.begin() as connection:
        connection.execute(sa.text('CREATE TABLE after_sale_refund_tracking (id INTEGER PRIMARY KEY, phase TEXT, last_error TEXT)'))
        connection.execute(sa.text("INSERT INTO after_sale_refund_tracking VALUES (1,'refunded','prior read failed')"))
        with Operations.context(MigrationContext.configure(connection)):migration['upgrade']()
        assert connection.execute(sa.text('SELECT phase,phase_evidence,note,last_error FROM after_sale_refund_tracking')).one()==('refunded',None,None,'prior read failed')
        with Operations.context(MigrationContext.configure(connection)):migration['downgrade']()
        assert connection.execute(sa.text('SELECT phase,last_error FROM after_sale_refund_tracking')).one()==('refunded','prior read failed')


def test_claim_history_keeps_current_node_evidence_and_precise_review_reason(tmp_path):
    for web,device,ids,_ in _e2e_setup(tmp_path):
        credential=ids['credential']
        heartbeat(device,credential,capabilities=AS_CAPABILITIES,client_version=CLIENT_VERSION)
        created=web.post('/v1/operation-runs/after-sale-claim',headers=CSRF,json={
            'executorId':ids['executorId'],'idempotencyKey':'synthetic-phase-claim-history',
            'items':[{'environmentSerial':'900001','orderNo':'ORDER'}]})
        assert created.status_code==202,created.text
        run=created.json()['data']['runId']
        task,lease=_lease_and_start(device,credential,expect_type='after.sale.claim.v1')
        receipt={'refundBillId':'12345','source':'existing','phase':'evidence_required',
                 'phaseLabel':'审核未通过 · 待补充凭证','phaseEvidence':evidence(),
                 'refundPath':'Cuenta original de pago','refundAccount':'****1234'}
        done=device.post(f'/v1/executor-channel/tasks/{task}/finish',headers=device_headers(credential),json={
            'leaseToken':lease,'outcome':'succeeded','resultCode':'existing-review-exception',
            'resultSummary':{'runStatus':'completed','phase':'after_sale.completed','totalCount':1,
                'progressTotal':1,'progressCompleted':1,'successCount':0,'failedCount':0,'skippedCount':1,
                'rows':[{'environmentSerial':'900001','orderNo':'ORDER','status':'blocked','refunds':[receipt]}]}})
        assert done.status_code==200,done.text
        for path in [f'/v1/operation-runs/after-sale-claim/{run}',f'/v1/operation-runs/after-sale-claim/history/{run}']:
            result=web.get(path).json()['data']['rows'][0]['refunds'][0]
            assert result['phaseEvidence']==evidence()
            assert '凭证不足' in result['detailsNote'] and '可补充凭证重新审核' in result['detailsNote']
        export=web.get(f'/v1/operation-runs/after-sale-claim/history/{run}/export')
        assert export.status_code==200
        sheet=load_workbook(io.BytesIO(export.content)).active
        assert '凭证不足' in sheet.cell(2,11).value


def test_submission_receipt_rejects_mismatched_current_node():
    from xynigo_auth.operation_contract import AfterSalePackageRefund
    with pytest.raises(ValidationError,match='node and phase'):
        AfterSalePackageRefund(refundBillId='12345',phase='processing',phaseEvidence=evidence())


@pytest.fixture
def tracking_store(monkeypatch):
    """Real SQL persistence; task payload decryption is outside these replay cases."""
    import uuid
    from sqlalchemy import create_engine
    from sqlalchemy.orm import Session
    from xynigo_auth.models import ExecutorTask
    from xynigo_auth.executor_service import ExecutorChannelService
    engine=create_engine('sqlite://')
    AfterSaleRefundTracking.__table__.create(engine)
    with Session(engine) as session:
        service=ExecutorChannelService(session)
        task=ExecutorTask(id=uuid.uuid4(),tenant_id=uuid.uuid4(),progress_summary={})
        monkeypatch.setattr(service,'_request_payload',lambda task: {'items':[
            {k:row()[k] for k in ('refundBillId','orderNo','environmentSerial')}]})
        yield service,session,task
    engine.dispose()


def observe(service, task, value):
    from datetime import datetime, timezone
    service._upsert_after_sale_track(task,{'rows':[value]},datetime.now(timezone.utc))
    return deepcopy(task.progress_summary['rows'][0])


def failed_row():
    return {k:row()[k] for k in ('refundBillId','orderNo','environmentSerial')} | {
        'status':'fail','errorSummary':'本次页面未完整加载'}


def reviewed_row():
    fresh=deepcopy(row())
    fresh.update(phase='reviewing',phaseLabel='审核中',note='正在重新审核',checkedAt='2026-01-02T10:00:00+00:00')
    fresh['phaseEvidence'].update(detail='En revisión',canSupplement=False,reason='',reasonSource='')
    return fresh


def test_legacy_unknown_environment_never_leaks_into_task_failure(tracking_store):
    service,session,task=tracking_store
    observe(service,task,row())
    record=session.scalar(select(AfterSaleRefundTracking))
    record.environment_serial=None
    record.refund_account='****9999'
    session.flush()
    task.progress_summary={}
    failed=observe(service,task,failed_row())
    assert not any(failed.get(key) for key in ('phase','phaseEvidence','refundAccount','checkedAt','note'))
    assert record.phase is None and record.refund_account is None


@pytest.mark.parametrize('had_prior_facts',[True,False])
def test_failure_replay_keeps_original_task_facts_and_leaves_newer_shared_observation(tracking_store,had_prior_facts):
    from xynigo_auth.models import ExecutorTask
    service,session,task_a=tracking_store
    if had_prior_facts:
        observe(service,task_a,row())
        task_a.progress_summary={}
    frozen=observe(service,task_a,failed_row())
    task_b=ExecutorTask(tenant_id=task_a.tenant_id,progress_summary={})
    observe(service,task_b,reviewed_row())
    assert observe(service,task_a,failed_row())==frozen
    record=session.scalar(select(AfterSaleRefundTracking))
    assert record.phase=='reviewing' and record.last_status=='ok' and not record.last_error
    assert record.phase_evidence['reason']==''
    if had_prior_facts:
        assert frozen['phase']=='evidence_required'
    else:
        assert frozen['phase']=='' and not frozen['checkedAt']


@pytest.mark.parametrize('status,label',[
    ('queued','等待回访'),('running','回访中'),('stopped','已停止'),('fail','本次状态未确认')])
def test_unconfirmed_export_distinguishes_current_execution_and_last_success(status,label):
    from xynigo_auth.after_sale_export import _row_values
    values=_row_values({**row(),'status':status,'countdown':'12:00:00'})
    assert values[6]==label+'；上次成功读取：审核未通过 · 待补充凭证'
    assert values[7]==''
    assert values[8]=='2026-01-01 10:00'
    assert '凭证不足' in values[9]
