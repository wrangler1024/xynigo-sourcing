"""Keep Dianxiaomi account labels without breaking existing import identities."""

import base64
from io import BytesIO

from openpyxl import Workbook, load_workbook
import pytest

from purchase_tool.procurement_import import ProcurementImportService
from purchase_tool.system_order_key import create_system_order_key
from test_procurement_import import source_workbook
from test_procurement_import_dedupe import edit_row, split_order_case, sync


@pytest.mark.parametrize('account,legacy_store,operator', [
    ('测试店铺-测试运营（二组）$', '测试店铺', '测试运营'),
    ('合成-品牌 - 合成运营 (三组) ￥', '合成-品牌', '合成运营'),
    ('合成店-合成运营（一组）¥', '合成店', '合成运营'),
    ('独立店铺  USD$', '独立店铺 USD', ''),
])
def test_preview_plan_and_export_keep_full_account_name(account, legacy_store, operator):
    original = load_workbook(BytesIO(source_workbook()))
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'order_'
    for index, values in enumerate(original.active.iter_rows(values_only=True)):
        values = list(values)
        if index:
            values[0] = account
        sheet.append(values)
    stream = BytesIO()
    workbook.save(stream)

    service = ProcurementImportService(sheet_gateway=object())
    result = service.parse('synthetic_store.xlsx', base64.b64encode(stream.getvalue()).decode())
    assert result['errorCount'] == 0
    assert all(row['store'] == account for row in result['preview'])
    plan = service.pending[result['planId']]
    for row in plan.rows:
        assert row.values['店铺'] == account
        assert row.values['运营'] == operator
        assert row.values['系统订单键'] == create_system_order_key(
            legacy_store, row.values['销售订单号'], row.values['包裹号'])
    exported, _, _ = service.export(result['planId'])
    exported_sheet = load_workbook(BytesIO(exported)).active
    assert [exported_sheet.cell(row, 4).value for row in (2, 3)] == [account, account]


def test_display_change_preserves_preexisting_partial_import_batch_and_key():
    service = ProcurementImportService(sheet_gateway=object())
    result = service.parse('synthetic_store.xlsx', base64.b64encode(source_workbook()).decode())
    # Captured from the deployed v0.17.17 parser with this synthetic fixture.
    assert result['importBatch'] == 'synthetic_store-e346ec6b351c'
    assert service.pending[result['planId']].rows[0].values['系统订单键'] == (
        'OK1-QR9RK-P2NH5-BM5B3-NHYT0')


@pytest.mark.parametrize('same_batch', [False, True])
@pytest.mark.parametrize('key_format', ['ok1', 'legacy'])
@pytest.mark.parametrize('split', [False, True])
def test_previous_short_store_rows_still_deduplicate(same_batch, key_format, split):
    service, gateway, plan = split_order_case(
        same_batch=same_batch, key_format=key_format, split=split)
    assert all(row.values['店铺'] == '测试店铺-测试运营（二组）$' for row in plan.rows)
    for index, (number, _) in enumerate(gateway.rows):
        edit_row(gateway, index, {'店铺': '测试店铺'})
        gateway.images[number] = True
    before = gateway.rows
    completed = sync(service, plan)
    assert completed['state'] == 'completed', completed
    assert completed['rowsWritten'] == 0 and completed['rowsExisting'] == 2
    assert gateway.rows == before
    assert gateway.append_calls == []
