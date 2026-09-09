"""Synthetic purchasing-authority and historical SKU compatibility cases."""
import base64
import json
from io import BytesIO
import unittest

from openpyxl import load_workbook

from purchase_tool.procurement_import import (
    ProcurementImportService, ProcurementImportError, require_importable,
)
from test_procurement_import import FakeSheetGateway, JPEG, source_workbook, xyp2_text, wait_for_sync
from test_procurement_import_validation import synthetic_workbook


def workbook_data(mutate):
    workbook = load_workbook(BytesIO(source_workbook()))
    sheet = workbook.active
    payload = json.loads(xyp2_text()[6:-7])
    mutate(sheet, payload)
    for index in range(2, sheet.max_row + 1):
        sheet.cell(index, 22).value = '[XYP2]' + json.dumps(payload) + '[/XYP2]'
    stream = BytesIO()
    workbook.save(stream)
    return base64.b64encode(stream.getvalue()).decode('ascii')


class ProcurementAuthorityTests(unittest.TestCase):
    def parse(self, mutate):
        service = ProcurementImportService(sheet_gateway=FakeSheetGateway(), sleep_fn=lambda _: None)
        self.service = service
        result = service.parse('synthetic.xlsx', workbook_data(mutate))
        return result, service.pending[result['planId']]

    def test_legacy_quantity_suffix_enriches_sales_without_changing_purchase(self):
        def mutate(sheet, payload):
            for index, item in enumerate(payload['i'], 2):
                item[0] = sheet.cell(index, 16).value + 'x' + str(sheet.cell(index, 20).value)
                item[4:6] = ['Purchase Color', 'Purchase Size']
        result, plan = self.parse(mutate)
        self.assertTrue(result['canImport'])
        self.assertEqual(result['warningCount'], 0)
        self.assertEqual([row.item_sales_amount for row in plan.rows], [150, 150])
        for row in plan.rows:
            self.assertEqual(row.values['主规格'], 'Purchase Color')
            self.assertEqual(row.values['次规格'], 'Purchase Size')
            self.assertIn('p=Purchase+Color', row.values['采购链接'])

    def test_real_sku_suffix_is_preserved_and_bad_quantity_suffix_is_not_guessed(self):
        for real_sku in (False, True):
            def mutate(sheet, payload):
                payload['i'][0][0] += 'x9'
                if real_sku:
                    sheet.cell(2, 16).value = payload['i'][0][0]
            result, plan = self.parse(mutate)
            self.assertTrue(result['canImport'])
            self.assertEqual(plan.rows[0].item_sales_amount, 150 if real_sku else None)
            if not real_sku:
                self.assertEqual(plan.rows[0].order_image, b'')

    def test_manual_detail_uses_purchase_link_without_repeating_sales_amount_or_image(self):
        def mutate(sheet, payload):
            manual = list(payload['i'][0])
            manual[0:3] = ['手工明细-1', '1234567', 'SYNTH-MANUAL']
            manual[4:6] = ['Multicolor', 'Set']
            manual[9] = 3
            payload['i'].append(manual)
        result, plan = self.parse(mutate)
        self.assertTrue(result['canImport'])
        self.assertEqual(result['detailCount'], 3)
        self.assertEqual(result['warningCount'], 1)
        self.assertTrue(result['preview'][-1]['sourceUnmatched'])
        manual = plan.rows[-1]
        self.assertEqual(manual.values['需求数量'], 3)
        self.assertEqual(manual.values['次规格'], 'Set')
        self.assertIn('goods_id=1234567', manual.values['采购链接'])
        self.assertIsNone(manual.item_sales_amount)
        self.assertIsNone(manual.values['销售订单金额'])
        self.assertEqual(manual.order_image, b'')
        self.assertEqual(manual.order_image_url, '')
        self.assertEqual(sum(row.values['销售订单金额'] is not None for row in plan.rows), 1)
        # openpyxl's synthetic workbook round-trip drops DXM cell-image ZIP
        # parts. Supply test images only for the two uniquely linked rows.
        for row in plan.rows[:-1]:
            row.order_image = JPEG
        self.service.validate_target(plan.plan_id,
            'https://tenant.feishu.cn/sheets/SheetToken123', 'sheetA')
        job = self.service.start_sheet_sync(plan.plan_id, confirm_write=True)
        status = wait_for_sync(self.service, job['jobId'])
        self.assertEqual(status['state'], 'completed', status)
        self.assertEqual(status['rowsWritten'], 3)
        self.assertEqual(status['skippedUnmatched'], 1)
        self.assertEqual(status['written'], 2)
        self.assertEqual(status['failed'], 0)

    def test_explicit_partial_import_keeps_failed_order_completely_out_and_all_issues(self):
        def invalid(order, variant, row):
            if order == 1 and variant == 1:
                row[21] = '[XYP2]{broken}[/XYP2]'
        service = ProcurementImportService()
        result = service.parse('synthetic.xlsx', synthetic_workbook(mutate=invalid),
                               allow_partial=True)
        self.assertTrue(result['partialImportSelected'])
        self.assertTrue(result['canImport'])
        self.assertEqual((result['successOrderCount'], result['failedOrderCount']), (2, 1))
        plan = service.pending[result['planId']]
        require_importable(plan)
        self.assertEqual(len(plan.issues), 1)
        self.assertNotIn('SYNTH-001', {row.values['销售订单号'] for row in plan.rows})
        exported, _, _ = service.export(plan.plan_id)
        self.assertEqual(load_workbook(BytesIO(exported)).active.max_row, 5)
        # A bad/missing order identity or an overlapping failed order still blocks.
        plan.issues.append({'level': 'error', 'orderNo': '', 'packageNo': ''})
        with self.assertRaises(ProcurementImportError):
            require_importable(plan)
        plan.issues[-1].update(orderNo='SYNTH-000', packageNo='SYNTH-PKG-000')
        with self.assertRaises(ProcurementImportError):
            require_importable(plan)

    def test_partial_option_requires_boolean_and_never_allows_no_valid_orders(self):
        service = ProcurementImportService()
        for value in ('true', 1, None):
            with self.assertRaises(ProcurementImportError):
                service.parse('synthetic.xlsx', synthetic_workbook(), allow_partial=value)
        def invalid(order, variant, row):
            row[21] = '[XYP2]{broken}[/XYP2]'
        with self.assertRaises(ProcurementImportError) as caught:
            service.parse('synthetic.xlsx', synthetic_workbook(mutate=invalid), allow_partial=True)
        self.assertFalse(caught.exception.diagnostics['canImport'])
