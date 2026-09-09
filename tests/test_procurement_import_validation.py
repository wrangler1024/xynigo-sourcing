"""Synthetic regression coverage for whole-order validation and diagnostics."""
import base64
import json
from io import BytesIO
import unittest

from openpyxl import Workbook, load_workbook

from purchase_tool.procurement_import import (
    CollaborationSheetTarget, ProcurementImportError, ProcurementImportService,
    parse_xyp2_remark,
)
from test_procurement_import import FakeSheetGateway, source_workbook, xyp2_text


def synthetic_workbook(count=3, mutate=None):
    source = load_workbook(BytesIO(source_workbook()))
    template = list(source.active.values)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'order_'
    sheet.append(template[0])
    for order in range(count):
        for variant, original in enumerate(template[1:]):
            row = list(original)
            row[1:3] = ['SYNTH-%03d' % order, 'SYNTH-PKG-%03d' % order]
            if mutate:
                mutate(order, variant, row)
            sheet.append(row)
    stream = BytesIO()
    workbook.save(stream)
    return base64.b64encode(stream.getvalue()).decode('ascii')


class ProcurementImportValidationTests(unittest.TestCase):
    def setUp(self):
        self.gateway = FakeSheetGateway()
        self.service = ProcurementImportService(sheet_gateway=self.gateway)

    def test_invalid_detail_excludes_whole_order_and_all_write_entrypoints_block(self):
        def invalid(order, variant, row):
            if order == 1 and variant == 1:
                row[19] = 'invalid'
        result = self.service.parse('synthetic.xlsx', synthetic_workbook(mutate=invalid))
        self.assertEqual((result['totalOrderCount'], result['successOrderCount'],
                          result['failedOrderCount'], result['detailCount']), (3, 2, 1, 4))
        self.assertFalse(result['canImport'])
        self.assertNotIn('SYNTH-001', {row['orderNo'] for row in result['preview']})
        self.assertEqual(result['issues'][0]['rowNumbers'], [5])
        self.assertEqual(result['issues'][0]['packageNo'], 'SYNTH-PKG-001')
        self.assertEqual(result['issues'][0]['field'], '单个产品数量')
        plan = self.service.pending[result['planId']]
        with self.assertRaisesRegex(ProcurementImportError, '阻断错误'):
            self.service.validate_target(plan.plan_id, 'synthetic', 'sheetA')
        with self.assertRaisesRegex(ProcurementImportError, '阻断错误'):
            self.service.export(plan.plan_id)
        # Simulate a legacy persisted plan with a previously validated target.
        plan.target = CollaborationSheetTarget('synthetic', 'sheetA', 'synthetic', 1)
        with self.assertRaisesRegex(ProcurementImportError, '阻断错误'):
            self.service.start_sheet_sync(plan.plan_id, confirm_write=True)
        self.assertFalse(self.service.sync_jobs)
        self.assertFalse(self.gateway.append_calls)
        self.assertFalse(self.gateway.presentation_calls)

    def test_all_failed_preserves_more_than_100_issues_and_all_order_counts(self):
        def invalid(order, variant, row):
            row[21] = '[XYP2]{broken}[/XYP2]'
        with self.assertRaises(ProcurementImportError) as caught:
            self.service.parse('synthetic.xlsx', synthetic_workbook(130, invalid))
        diagnostics = caught.exception.diagnostics
        self.assertEqual((diagnostics['totalOrderCount'], diagnostics['successOrderCount'],
                          diagnostics['failedOrderCount']), (130, 0, 130))
        self.assertEqual(len(diagnostics['issues']), 260)
        self.assertEqual(diagnostics['issues'][-1]['rowNumbers'], [261])
        self.assertFalse(diagnostics['canImport'])
        self.assertFalse(self.service.pending)

    def test_partial_failure_preserves_more_than_100_issues(self):
        def invalid(order, variant, row):
            if order:
                row[21] = '[XYP2]{broken}[/XYP2]'
        result = self.service.parse('synthetic.xlsx', synthetic_workbook(60, invalid))
        self.assertEqual(len(result['issues']), 118)
        self.assertEqual(result['failedOrderCount'], 59)
        self.assertEqual(result['successOrderCount'], 1)

    def test_damaged_sibling_remark_blocks_whole_order(self):
        def invalid(order, variant, row):
            if order == 1 and variant == 1:
                row[21] = xyp2_text().replace('[/XYP2]', '')
        result = self.service.parse('synthetic.xlsx', synthetic_workbook(mutate=invalid))
        self.assertEqual(result['detailCount'], 4)
        self.assertEqual(result['issues'][0]['rowNumbers'], [5])
        self.assertFalse(result['canImport'])

    def test_absent_sibling_remark_warns_without_blocking_order_level_metadata(self):
        def missing(order, variant, row):
            if variant:
                row[21] = '普通人工备注'
        result = self.service.parse('synthetic.xlsx', synthetic_workbook(1, missing))
        self.assertTrue(result['canImport'])
        self.assertEqual(result['detailCount'], 2)
        self.assertEqual(result['issues'][0]['code'], 'xyp2_inherited')

    def test_duplicate_blocks_compare_meaning_and_conflicting_or_damaged_blocks_fail(self):
        payload = json.loads(xyp2_text()[6:-7])
        equivalent = '[XYP2]' + json.dumps(payload, sort_keys=True, indent=2) + '[/XYP2]'
        self.assertEqual(parse_xyp2_remark(xyp2_text() + equivalent).block_count, 2)
        payload['i'][1][8] += 1
        conflict = '[XYP2]' + json.dumps(payload) + '[/XYP2]'
        for text in [xyp2_text() + conflict, xyp2_text() + '[XYP2]{',
                     xyp2_text() + '[XYP2]{broken}[/XYP2]']:
            with self.subTest(text=text), self.assertRaises(ProcurementImportError):
                parse_xyp2_remark(text)
        def duplicate(order, variant, row):
            row[21] = xyp2_text() + equivalent
        result = self.service.parse('synthetic.xlsx', synthetic_workbook(1, duplicate))
        self.assertTrue(result['canImport'])
        self.assertEqual(result['detailCount'], 2)
        self.assertEqual(result['warningCount'], 2)

    def test_different_valid_remarks_conflict_across_rows(self):
        def conflicting(order, variant, row):
            if order == 1 and variant:
                payload = json.loads(xyp2_text()[6:-7])
                payload['i'][1][8] += 1
                row[21] = '[XYP2]' + json.dumps(payload) + '[/XYP2]'
        result = self.service.parse('synthetic.xlsx', synthetic_workbook(mutate=conflicting))
        self.assertEqual(result['failedOrderCount'], 1)
        self.assertEqual(result['issues'][0]['rowNumbers'], [4, 5])
        self.assertEqual(result['issues'][0]['code'], 'xyp2_conflict')

    def test_quantity_difference_remains_warning_and_fixed_input_can_import(self):
        def differing(order, variant, row):
            row[19] = 2
        result = self.service.parse('synthetic.xlsx', synthetic_workbook(1, differing))
        self.assertTrue(result['canImport'])
        self.assertEqual(result['warningCount'], 2)
        self.assertEqual(result['quantityCount'], 2)
        self.assertEqual(result['issues'][0]['code'], 'quantity_mismatch')
        fixed = self.service.parse('synthetic.xlsx', synthetic_workbook())
        self.assertTrue(fixed['canImport'])
        self.assertEqual(fixed['failedOrderCount'], 0)

    def test_unmatched_source_is_not_assigned_by_row_order(self):
        def unmatched(order, variant, row):
            if order == 1:
                row[15] = 'UNKNOWN-SKU-%d' % variant
                row[17] = 'UNKNOWN-SPEC-%d' % variant
        result = self.service.parse('synthetic.xlsx', synthetic_workbook(mutate=unmatched))
        self.assertEqual(result['detailCount'], 6)
        self.assertEqual(result['issues'][0]['code'], 'source_match_warning')
        self.assertTrue(result['canImport'])
        unmatched = [row for row in self.service.pending[result['planId']].rows
                     if row.values['销售订单号'] == 'SYNTH-001']
        self.assertTrue(all(row.item_sales_amount is None and not row.order_image
                            and not row.order_image_url for row in unmatched))

    def test_invalid_numeric_rows_and_missing_identity_are_all_reported(self):
        def invalid(order, variant, row):
            if order == 1:
                row[18] = 'nan' if variant else 'inf'
                row[19] = -1
            if order == 2:
                row[2] = ''
        result = self.service.parse('synthetic.xlsx', synthetic_workbook(mutate=invalid))
        self.assertEqual(result['successOrderCount'], 1)
        self.assertEqual(len(result['issues']), 6)
        self.assertTrue(all(item['rowNumbers'] for item in result['issues']))


if __name__ == '__main__':
    unittest.main()
